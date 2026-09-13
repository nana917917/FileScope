"""PowerPoint extraction: python-pptx plus comments/charts/SmartArt."""

from __future__ import annotations

from ..core.models import ChunkKind, FileEntry
from ..core.paths import POWERPOINT_EXTS
from ..errors import Issue, Severity
from .base import ExtractContext, ExtractOptions, ExtractResult, Sink
from .ooxml import iter_parts, open_zip, xml_texts


class PowerPointExtractor:
    name = "powerpoint"
    version = "5.0"

    def supports(self, entry: FileEntry, options: ExtractOptions) -> bool:
        return entry.extension in POWERPOINT_EXTS

    def extract(self, context: ExtractContext, sink: Sink) -> ExtractResult:
        result = ExtractResult()
        limit = context.options.limits.office_max_bytes
        if context.entry.size > limit:
            result.skipped = True
            result.reason = f"Office上限（{limit // (1024 * 1024)}MB）を超えています"
            return result

        try:
            from pptx import Presentation
        except ImportError:
            result.skipped = True
            result.reason = "python-pptx が未導入です"
            result.warnings.append(
                Issue(path=context.entry.path, code="ppt-dependency", message=result.reason, severity=Severity.WARN)
            )
            return result

        try:
            presentation = Presentation(context.path)
        except Exception as exc:
            result.skipped = True
            result.reason = f"開けませんでした: {type(exc).__name__}"
            result.warnings.append(
                Issue(path=context.entry.path, code="ppt-open", message=result.reason, severity=Severity.WARN)
            )
            return result

        for index, slide in enumerate(presentation.slides, 1):
            hidden = str(slide._element.get("show", "1")) == "0"
            prefix = f"スライド{index}" + ("（非表示）" if hidden else "")
            for shape in slide.shapes:
                sink.text(getattr(shape, "name", ""), ChunkKind.INTERNAL, f"{prefix} 図形名")
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        sink.text(paragraph.text, ChunkKind.SLIDE, prefix)
                if getattr(shape, "has_table", False):
                    for row_index, row in enumerate(shape.table.rows, 1):
                        for column_index, cell in enumerate(row.cells, 1):
                            sink.text(
                                cell.text,
                                ChunkKind.SLIDE,
                                f"{prefix} 表{row_index}行{column_index}列",
                            )
                if getattr(shape, "has_chart", False):
                    try:
                        chart = shape.chart
                        sink.text(chart.chart_title.text_frame.text, ChunkKind.CHART, prefix)
                        for series in chart.plots[0].series:
                            sink.text(str(series.name), ChunkKind.CHART, f"{prefix} 系列名")
                    except Exception:
                        sink.text("", ChunkKind.CHART, prefix)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                sink.text(notes, ChunkKind.NOTES, f"{prefix} ノート")

        self._extract_parts(context, sink, result)
        return result

    def _extract_parts(self, context: ExtractContext, sink: Sink, result: ExtractResult) -> None:
        archive = open_zip(context.path)
        if archive is None:
            return
        with archive:
            for _name, data in iter_parts(
                archive,
                lambda member: "/comments/" in member and member.endswith(".xml"),
                max_parts=80,
            ):
                for value in xml_texts(data, {"t"}):
                    sink.text(value, ChunkKind.COMMENT, "コメント")

            for _name, data in iter_parts(
                archive,
                lambda member: member.startswith("ppt/charts/") and member.endswith(".xml"),
                max_parts=80,
            ):
                for value in xml_texts(data, {"t", "v"}):
                    sink.text(value, ChunkKind.CHART, "グラフ")

            for _name, data in iter_parts(
                archive,
                lambda member: member.startswith("ppt/diagrams/") and member.endswith(".xml"),
                max_parts=80,
            ):
                for value in xml_texts(data, {"t"}):
                    sink.text(value, ChunkKind.TEXTBOX, "SmartArt")
