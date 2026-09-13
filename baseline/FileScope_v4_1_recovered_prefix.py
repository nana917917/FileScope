from __future__ import annotations

import csv
import ctypes
import datetime as dt
import io
import itertools
import json
import os
import queue
import re
import shutil
import subprocess
import tempfile
import threading
import time
import traceback
import unicodedata
from collections import defaultdict
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# -----------------------------------------------------------------------------
# Optional document readers. The program itself never sends file contents or
# search terms to the network. Network access only happens indirectly when the
# selected path itself is a network drive / SharePoint-OneDrive synced folder.
# -----------------------------------------------------------------------------
try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    load_workbook = None
    get_column_letter = None

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    from pyxlsb import open_workbook as open_xlsb_workbook
except ImportError:
    open_xlsb_workbook = None

try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    import win32com.client as win32
except ImportError:
    win32 = None


APP_NAME = "FileScope"
EXCEL_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".xlsb"}
PDF_EXTS = {".pdf"}
WORD_EXTS = {".docx", ".docm"}
POWERPOINT_EXTS = {".pptx", ".pptm"}
TEXT_EXTS = {
    ".txt", ".csv", ".tsv", ".log", ".md", ".ini", ".json", ".xml",
    ".html", ".htm", ".py", ".bat", ".ps1", ".sql", ".yaml", ".yml",
    ".cfg", ".conf", ".properties", ".rtf",
}
ALL_SUPPORTED_EXTS = EXCEL_EXTS | PDF_EXTS | WORD_EXTS | POWERPOINT_EXTS | TEXT_EXTS
TEXT_PROBE_BYTES = 64 * 1024
UNKNOWN_TEXT_MAX_BYTES = 50 * 1024 * 1024
OBVIOUS_BINARY_EXTS = {
    ".exe", ".dll", ".bin", ".iso", ".zip", ".7z", ".rar", ".tar", ".gz", ".png", ".jpg", ".jpeg",
    ".gif", ".bmp", ".tif", ".tiff", ".webp", ".mp3", ".wav", ".mp4", ".mov", ".avi", ".db", ".sqlite",
}
HYPHEN_CHARS = "-‐‑‒–—―−ーｰ－"

DEFAULT_WORKERS = 4
INITIAL_DISPLAY_ROWS = 500
LOAD_MORE_ROWS = 500
KEYWORD_DETAIL_ROWS = 1000
UI_REFRESH_MS = 150
AUTOSAVE_FLUSH_EVERY = 100
OCR_PAGE_TIMEOUT_SECONDS = 30

# Safety limits for PCs with relatively little free space.
# Remote Office/PDF files larger than this are parsed in-place rather than
# duplicated into the system TEMP directory. With four workers this caps the
# tool-created staging footprint to about 1 GiB in the usual worst case.
MAX_REMOTE_STAGE_BYTES = 256 * 1024 * 1024
TEMP_FREE_RESERVE_BYTES = 4 * 1024 * 1024 * 1024
AUTOSAVE_LOW_DISK_STOP_BYTES = 3 * 1024 * 1024 * 1024

# Keep the UI/RAM bounded even for pathological queries (for example a common
# word that matches almost every line/cell). Every hit is still streamed to the
# autosave CSV; only the in-memory interactive result set is capped.
MAX_IN_MEMORY_HITS = 250_000


class SearchStopped(Exception):
    pass


@dataclass(frozen=True)
class SearchOptions:
    roots: tuple[str, ...]
    query: str
    exclude_query: str
    include_subfolders: bool
    include_excel: bool
    include_pdf: bool
    include_word: bool
    include_powerpoint: bool
    include_text: bool
    include_unknown_text: bool
    pdf_ocr_mode: str
    case_sensitive: bool
    ignore_width: bool
    part_number_mode: bool
    search_formula: bool
    search_mode: str
    search_path_names: bool
    stage_remote_files: bool
    workers: int = DEFAULT_WORKERS


@dataclass
class SearchHit:
    file_type: str
    path: str
    place: str
    cell: str
    value: str
    matched_keywords: tuple[str, ...] = field(default_factory=tuple)
    severity: str = "HIT"  # HIT / INFO / ERROR


@dataclass
class RuntimeStats:
    total_files: int = 0
    completed_files: int = 0
    hit_count: int = 0
    hit_files: int = 0
    errors: int = 0
    started_at: float = 0.0
    finished_at: float = 0.0
    phase: str = "待機中"
    last_file: str = ""
    active_files: dict[str, tuple[str, float]] = field(default_factory=dict)
    file_type_seconds: dict[str, float] = field(default_factory=lambda: defaultdict(float))


# -----------------------------------------------------------------------------
# Search matcher
# -----------------------------------------------------------------------------
class SearchMatcher:
    """Immutable matcher for one search run.

    Part-number mode intentionally treats hyphen variants and whitespace as
    equivalent, and supports '*' as a wildcard. Normal substring matching is
    always tried first so the special mode never makes ordinary searches worse.
    """

    def __init__(
        self,
        query: str,
        exclude_query: str,
        *,
        case_sensitive: bool,
        ignore_width: bool,
        part_number_mode: bool,
        search_mode: str,
    ):
        self.query = query.strip()
        self.exclude_query = exclude_query.strip()
        self.case_sensitive = case_sensitive
        self.ignore_width = ignore_width
        self.part_number_mode = part_number_mode
        self.search_mode = "AND" if str(search_mode).upper() == "AND" else "OR"

        # Threshold syntax for "N of these terms" searches, e.g.
        #   2of(AAA,落下,耐久)
        # This is intentionally file-level: the matching terms may live in
        # different cells, sheets, pages, slides, or lines of the same file.
        self.threshold_min: int | None = None
        self.threshold_terms_raw: list[str] = []
        threshold = re.fullmatch(r"\s*(\d+)\s*of\s*\((.*)\)\s*", self.query, flags=re.IGNORECASE)
        if threshold:
            raw_terms = [t.strip() for t in re.split(r"[,、;；]+", threshold.group(2)) if t.strip()]
            requested = int(threshold.group(1))
            if raw_terms and 1 <= requested <= len(raw_terms):
                self.threshold_min = requested
                self.threshold_terms_raw = raw_terms
                self.explicit_ops = True
                self.groups_raw = [[term] for term in raw_terms]
            else:
                self.explicit_ops = any(op in self.query for op in [",", "、", "&", ";", "＆"])
                self.groups_raw = self._parse_query(self.query)
        else:
            self.explicit_ops = any(op in self.query for op in [",", "、", "&", ";", "＆"])
            self.groups_raw = self._parse_query(self.query)

        if self.threshold_min is None and not self.explicit_ops and self.groups_raw:
            flat = [term for group in self.groups_raw for term in group]
            if self.search_mode == "OR":
                self.groups_raw = [[term] for term in flat]
            else:
                self.groups_raw = [flat]

        self.groups = [
            [self._prepare_term(term) for term in group]
            for group in self.groups_raw
            if group
        ]
        self.exclude_terms = [
            self._prepare_term(term)
            for term in self._split_plain_terms(self.exclude_query)
            if term
        ]
        self.display_names = [
            group[0] if len(group) == 1 else "&".join(group)
            for group in self.groups_raw
        ]
        self.threshold_display = (
            f"{self.threshold_min}of({','.join(self.threshold_terms_raw)})"
            if self.threshold_min is not None else ""
        )

    def _normalize(self, text: object) -> str:
        if text is None:
            return ""
        value = str(text)
        if self.ignore_width:
            value = unicodedata.normalize("NFKC", value)
        if not self.case_sensitive:
            value = value.casefold()
        return value

    def _canonical_part(self, text: object, *, keep_star: bool = False) -> str:
        if text is None:
            return ""
        value = unicodedata.normalize("NFKC", str(text))
        if not self.case_sensitive:
            value = value.casefold()

        out: list[str] = []
        for ch in value:
            if ch in HYPHEN_CHARS or ch.isspace():
                continue
            if keep_star and ch == "*":
                out.append(ch)
                continue
            if ch.isalnum():
                out.append(ch)
        return "".join(out)

    def _prepare_term(self, raw: str) -> tuple[str, str, re.Pattern[str] | None]:
        norm = self._normalize(raw)
        part = self._canonical_part(raw, keep_star=True) if self.part_number_mode else ""
        wildcard_re = None
        if self.part_number_mode and "*" in part:
            wildcard_re = re.compile(re.escape(part).replace(r"\*", ".*"))
        return norm, part, wildcard_re

    @staticmethod
    def _split_plain_terms(query: str) -> list[str]:
        query = query.strip().replace("　", " ")
        if not query:
            return []
        return [term for term in re.split(r"\s+", query) if term]

    def _parse_query(self, query: str) -> list[list[str]]:
        query = query.strip().replace("　", " ")
        if not query:
            return []

        replacements = {
            "＆": "&", "，": ",", "、": ",", "；": ";", "｜": "|", "／": "/",
        }
        for old, new in replacements.items():
            query = query.replace(old, new)

        query = re.sub(r"\b(かつ|且つ|and)\b", "&", query, flags=re.IGNORECASE)
        query = re.sub(r"\b(または|又は|もしくは|or)\b", ",", query, flags=re.IGNORECASE)

        if any(op in query for op in [",", ";", "&"]):
            groups: list[list[str]] = []
            for clause in re.split(r"[;,]+", query):
                clause = clause.strip()
                if not clause:
                    continue
                if "&" in clause:
                    terms: list[str] = []
                    for part in clause.split("&"):
                        terms.extend(self._split_plain_terms(part))
                else:
                    # Historical behavior: spaces inside an explicit OR clause are AND.
                    terms = self._split_plain_terms(clause)
                if terms:
                    groups.append(terms)
            return groups

        terms = self._split_plain_terms(query)
        return [terms] if terms else []

    def _make_targets(self, candidates: list[object] | tuple[object, ...]) -> tuple[str, str]:
        values = [c for c in candidates if c is not None and str(c) != ""]
        target_norm = "\n".join(self._normalize(c) for c in values)
        target_part = ""
        if self.part_number_mode:
            target_part = "\n".join(self._canonical_part(c) for c in values)
        return target_norm, target_part

    @staticmethod
    def _term_matches(prepared: tuple[str, str, re.Pattern[str] | None], target_norm: str, target_part: str) -> bool:
        norm, part, wildcard_re = prepared
        if norm and norm in target_norm:
            return True
        if not part:
            return False
        if wildcard_re is not None:
            return wildcard_re.search(target_part) is not None
        return part in target_part

    def evaluate(self, candidates: list[object] | tuple[object, ...]) -> tuple[str, ...]:
        if not self.groups:
            return ()
        target_norm, target_part = self._make_targets(candidates)
        if not target_norm and not target_part:
            return ()

        for term in self.exclude_terms:
            if self._term_matches(term, target_norm, target_part):
                return ()

        if self.threshold_min is not None:
            count = sum(
                1 for group in self.groups
                if group and self._term_matches(group[0], target_norm, target_part)
            )
            return (self.threshold_display,) if count >= self.threshold_min else ()

        matched: list[str] = []
        for display, group in zip(self.display_names, self.groups):
            if all(self._term_matches(term, target_norm, target_part) for term in group):
                matched.append(display)
        return tuple(matched)

    def unit_terms(self, candidates: list[object] | tuple[object, ...]) -> tuple[str, ...]:
        """Return indiMilual search terms that match this unit. Used for file-level AND."""
        if not self.groups:
            return ()
        target_norm, target_part = self._make_targets(candidates)
        if not target_norm and not target_part:
            return ()
        for term in self.exclude_terms:
            if self._term_matches(term, target_norm, target_part):
                return ()
        matched: list[str] = []
        seen: set[str] = set()
        for group_raw, group in zip(self.groups_raw, self.groups):
            for raw, prepared in zip(group_raw, group):
                if raw not in seen and self._term_matches(prepared, target_norm, target_part):
                    matched.append(raw)
                    seen.add(raw)
        return tuple(matched)

    def file_displays(self, present_terms: set[str]) -> tuple[str, ...]:
        """Return query expressions satisfied somewhere in the same file."""
        if self.threshold_min is not None:
            present = sum(1 for raw in self.threshold_terms_raw if raw in present_terms)
            return (self.threshold_display,) if present >= self.threshold_min else ()
        matched: list[str] = []
        for display, group_raw in zip(self.display_names, self.groups_raw):
            if all(raw in present_terms for raw in group_raw):
                matched.append(display)
        return tuple(matched)

    def terms_for_displays(self, displays: tuple[str, ...]) -> set[str]:
        if self.threshold_min is not None and self.threshold_display in set(displays):
            return set(self.threshold_terms_raw)
        wanted: set[str] = set()
        display_set = set(displays)
        for display, group_raw in zip(self.display_names, self.groups_raw):
            if display in display_set:
                wanted.update(group_raw)
        return wanted

    def contains_excluded(self, candidates: list[object] | tuple[object, ...]) -> bool:
        if not self.exclude_terms:
            return False
        target_norm, target_part = self._make_targets(candidates)
        return any(self._term_matches(term, target_norm, target_part) for term in self.exclude_terms)

    def snippet(self, text: object, max_len: int = 260) -> str:
        value = " ".join(str(text or "").replace("\r", "\n").split())
        if len(value) <= max_len:
            return value

        norm = self._normalize(value)
        positions: list[int] = []
        for group in self.groups:
            for term, _, _ in group:
                if term:
                    pos = norm.find(term)
                    if pos >= 0:
                        positions.append(pos)
        pos = min(positions) if positions else 0
        half = max_len // 2
        start = max(0, pos - half)
        end = min(len(value), start + max_len)
        prefix = "..." if start else ""
        suffix = "..." if end < len(value) else ""
        return prefix + value[start:end] + suffix


# -----------------------------------------------------------------------------
# File helpers
# -----------------------------------------------------------------------------
def get_config_path() -> str:
    if os.name == "nt":
        base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    else:
        base = os.path.join(os.path.expanduser("~"), ".config")
    folder = os.path.join(base, APP_NAME)
    try:
        os.makedirs(folder, exist_ok=True)
        return os.path.join(folder, "settings.json")
    except Exception:
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), "search_tool_settings_v2.json")


CONFIG_PATH = get_config_path()
LEGACY_CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "search_tool_settings.json")


def is_remote_path(path: str) -> bool:
    """Return True for UNC or Windows mapped network drives."""
    try:
        path = os.path.abspath(path)
    except Exception:
        pass
    if str(path).startswith("\\\\"):
        return True
    if os.name != "nt":
        return False
    try:
        drive, _ = os.path.splitdrive(path)
        if not drive:
            return False
        DRIVE_REMOTE = 4
        return ctypes.windll.kernel32.GetDriveTypeW(drive + "\\") == DRIVE_REMOTE
    except Exception:
        return False


def find_tesseract_executable() -> str | None:
    """Find Tesseract without requiring users to edit PATH manually."""
    if pytesseract is None:
        return None
    candidates: list[str] = []
    env_cmd = os.environ.get("TESSERACT_CMD", "").strip()
    if env_cmd:
        candidates.append(env_cmd)
    which = shutil.which("tesseract")
    if which:
        candidates.append(which)
    app_dir = os.path.dirname(os.path.abspath(__file__))
    candidates.extend([
        os.path.join(app_dir, "Tesseract-OCR", "tesseract.exe"),
        os.path.join(app_dir, "tesseract", "tesseract.exe"),
        os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"), "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"), "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Tesseract-OCR", "tesseract.exe"),
    ])
    for candidate in candidates:
        if candidate and os.path.isfile(candidate):
            try:
                pytesseract.pytesseract.tesseract_cmd = candidate
                return candidate
            except Exception:
                pass
    return None


def detect_ocr_capability() -> tuple[bool, str, str]:
    """Return (ready, language_expression, human_status)."""
    if pdfium is None:
        return False, "", "OCR未導入: pypdfium2 が必要"
    if pytesseract is None:
        return False, "", "OCR未導入: pytesseract が必要"
    cmd = find_tesseract_executable()
    if not cmd:
        return False, "", "OCR未導入: Tesseract本体が見つかりません"
    try:
        languages = set(pytesseract.get_languages(config=""))
    except Exception as exc:
        return False, "", f"OCR確認エラー: {type(exc).__name__}"
    if "jpn" in languages and "eng" in languages:
        return True, "jpn+eng", "OCR利用可: 日本語+英語"
    if "jpn" in languages:
        return True, "jpn", "OCR利用可: 日本語（eng未導入）"
    if "eng" in languages:
        return True, "eng", "OCR利用可: 英語のみ（jpn未導入）"
    return False, "", "OCR未導入: jpn/eng 言語データがありません"


@contextmanager
def local_parse_path(path: str, options: SearchOptions):
    """For remote structured documents, copy once then parse locally.

    XLSX/DOCX/PPTX are ZIP containers. Parsing them directly on SMB can cause
    many small network reads. A sequential copy to a local temp file is usually
    more stable and often much faster. Original files are never modified.
    """
    ext = os.path.splitext(path)[1].lower()
    should_stage = options.stage_remote_files and is_remote_path(path) and ext not in TEXT_EXTS
    if not should_stage:
        yield path
        return

    # Protect small system drives. A very large online/SMB file is safer to
    # parse directly than to create another full copy on C:.
    try:
        source_size = os.path.getsize(path)
    except Exception:
        source_size = 0

    if source_size > MAX_REMOTE_STAGE_BYTES:
        yield path
        return

    try:
        temp_dir = tempfile.gettempdir()
        free_bytes = shutil.disk_usage(temp_dir).free
        required = max(source_size, 1) + TEMP_FREE_RESERVE_BYTES
        if free_bytes < required:
            yield path
            return
    except Exception:
        pass

    temp_path = ""
    try:
        fd, temp_path = tempfile.mkstemp(prefix="file_search_v2_", suffix=ext)
        os.close(fd)
        shutil.copyfile(path, temp_path)
        yield temp_path
    finally:
        if temp_path:
            try:
                os.remove(temp_path)
            except Exception:
                pass


def excel_cell_candidates(value: object, number_format: str = "") -> list[str]:
    if value is None:
        return []

    candidates = [str(value)]

    if isinstance(value, bool):
        candidates.extend(["TRUE" if value else "FALSE", "true" if value else "false"])
    elif isinstance(value, int):
        candidates.extend([
            f"{value}.0", f"{value}.00", f"{value:03d}", f"{value:04d}", f"{value:05d}",
        ])
    elif isinstance(value, float):
        if value.is_integer():
            iv = int(value)
            candidates.extend([
                str(iv), f"{iv}.0", f"{iv}.00", f"{iv:03d}", f"{iv:04d}", f"{iv:05d}",
            ])
        decimals = decimal_places_from_format(number_format)
        if decimals is not None:
            candidates.append(f"{value:.{decimals}f}")
    elif isinstance(value, (dt.datetime, dt.date)):
        candidates.extend([
            value.strftime("%Y/%m/%d"), value.strftime("%Y-%m-%d"),
            f"{value.year}/{value.month}/{value.day}",
        ])
        if isinstance(value, dt.datetime):
            candidates.extend([
                value.strftime("%Y/%m/%d %H:%M"), value.strftime("%Y/%m/%d %H:%M:%S"),
            ])

    return list(dict.fromkeys(candidates))


def decimal_places_from_format(fmt: str) -> int | None:
    if not fmt or "." not in str(fmt):
        return None
    after = str(fmt).split(".", 1)[1]
    count = 0
    for ch in after:
        if ch in ("0", "#"):
            count += 1
        else:
            break
    return min(count, 10) if count else None



def looks_like_text_file(path: str) -> bool:
    """Lightweight local-only probe for proprietary text-like files (.td2/.dat/.prm etc.)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in OBVIOUS_BINARY_EXTS:
        return False
    try:
        size = os.path.getsize(path)
        if size > UNKNOWN_TEXT_MAX_BYTES:
            return False
        with urn PS   for c_l-onlxcepDATA      _rogue
        with    if self._term_matches(term   ES:
            return Fa�]
   lc