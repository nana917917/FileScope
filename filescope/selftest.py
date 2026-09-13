"""Release self test (spec section 5 of the RC brief).

Runs entirely inside a temporary folder and never reads the user's documents.
Every check reports PASS / WARN / SKIP / FAIL, so a machine without Tesseract
still produces a useful result instead of a failure.
"""

from __future__ import annotations

import os
import shutil
import sqlite3
import tempfile
import threading
import time
from dataclasses import dataclass, field

from . import diagnostics, paths
from .config import SCHEMA_VERSION, Settings
from .core.coordinator import SearchConfig, SearchSession
from .index.database import IndexDatabase, extraction_fingerprint
from .platform import tesseract as tesseract_module
from .version import __version__

PASS = "PASS"
WARN = "WARN"
SKIP = "SKIP"
FAIL = "FAIL"


@dataclass
class Check:
    name: str
    status: str
    detail: str = ""


@dataclass
class Report:
    checks: list[Check] = field(default_factory=list)

    def add(self, name: str, status: str, detail: str = "") -> None:
        self.checks.append(Check(name, status, detail))

    def count(self, status: str) -> int:
        return sum(1 for check in self.checks if check.status == status)

    @property
    def failed(self) -> bool:
        return self.count(FAIL) > 0

    def as_text(self) -> str:
        width = max((len(check.name) for check in self.checks), default=10)
        lines = [f"FileScope self test {__version__}", ""]
        for check in self.checks:
            line = f"{check.status:4}  {check.name.ljust(width)}"
            if check.detail:
                line += f"  {check.detail}"
            lines.append(line)
        lines.append("")
        lines.append(
            f"summary: PASS {self.count(PASS)} / WARN {self.count(WARN)}"
            f" / SKIP {self.count(SKIP)} / FAIL {self.count(FAIL)}"
        )
        if self.count(SKIP):
            lines.append("SKIP は環境依存（例: Tesseract未導入）で、通常検索には影響しません。")
        return "\n".join(lines)


def run(*, keep_workspace: bool = False, quiet: bool = False) -> Report:
    report = Report()
    workspace = tempfile.mkdtemp(prefix="filescope-selftest-")
    try:
        _check_environment(report)
        _check_settings(report, workspace)
        _check_sqlite(report, workspace)
        corpus = _build_corpus(report, workspace)
        _check_search(report, corpus, workspace)
        _check_index_equivalence(report, corpus, workspace)
        _check_staging(report, workspace)
        _check_ocr(report, workspace)
    except Exception as exc:  # a broken selftest must still explain itself
        report.add("self test harness", FAIL, f"{type(exc).__name__}: {exc}")
    finally:
        if keep_workspace:
            report.add("workspace", WARN, f"kept: {workspace}")
        else:
            try:
                shutil.rmtree(workspace, ignore_errors=True)
                report.add("workspace cleanup", PASS if not os.path.exists(workspace) else WARN, workspace)
            except OSError as exc:
                report.add("workspace cleanup", WARN, str(exc))
    _ = quiet
    return report


def _check_environment(report: Report) -> None:
    info = diagnostics.collect()
    rows = dict(info.rows)
    report.add("python", PASS, rows.get("Python", "?"))
    report.add("sqlite", PASS, rows.get("SQLite", "?"))
    report.add("fts5", PASS if rows.get("FTS5") == "OK" else FAIL, rows.get("FTS5", "?"))
    report.add(
        "fts5 trigram",
        PASS if rows.get("FTS5 trigram") == "OK" else WARN,
        rows.get("FTS5 trigram", "?"),
    )
    for name in ("openpyxl", "pypdf", "docx", "pptx"):
        report.add(name, PASS if rows.get(name, "").startswith("OK") else WARN, rows.get(name, "?"))


def _check_settings(report: Report, workspace: str) -> None:
    path = os.path.join(workspace, "settings.json")
    settings = Settings()
    settings.defaults.query = "セルフテスト"
    settings.index_max_bytes = 32 * 1024 * 1024
    ok = settings.save(path)
    loaded = Settings.load(path)
    report.add(
        "settings write/read",
        PASS if ok and loaded.defaults.query == "セルフテスト" else FAIL,
        path,
    )
    with open(path, "w", encoding="utf-8") as handle:
        handle.write("{ broken")
    recovered = Settings.load(path)
    report.add(
        "settings corruption fallback",
        PASS if recovered.schema_version == SCHEMA_VERSION else FAIL,
        "defaults restored",
    )
    # A real V4 file must migrate.
    v4_path = os.path.join(workspace, "settings_v4.json")
    with open(v4_path, "w", encoding="utf-8") as handle:
        handle.write('{"folder": "C:/x", "keyword": "AAA", "pdf_ocr_mode": "全ページ"}')
    migrated = Settings.load(v4_path)
    report.add(
        "V4 settings migration",
        PASS if migrated.defaults.root == "C:/x" and migrated.defaults.pdf_ocr_mode == "all" else FAIL,
        f"root={migrated.defaults.root} ocr={migrated.defaults.pdf_ocr_mode}",
    )


def _check_sqlite(report: Report, workspace: str) -> None:
    path = os.path.join(workspace, "probe.sqlite3")
    try:
        connection = sqlite3.connect(path)
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("CREATE TABLE t(x TEXT)")
        connection.execute("INSERT INTO t VALUES ('ok')")
        connection.commit()
        value = connection.execute("SELECT x FROM t").fetchone()[0]
        connection.close()
        report.add("sqlite write", PASS if value == "ok" else FAIL, path)
    except sqlite3.Error as exc:
        report.add("sqlite write", FAIL, f"{type(exc).__name__}: {exc}")


def _build_corpus(report: Report, workspace: str) -> str:
    corpus = os.path.join(workspace, "corpus")
    os.makedirs(corpus, exist_ok=True)
    with open(os.path.join(corpus, "a.txt"), "w", encoding="utf-8") as handle:
        handle.write("AAA 耐久 試験\nBBB 電源\n")
    with open(os.path.join(corpus, "b.txt"), "w", encoding="utf-8") as handle:
        handle.write("BBB CCC\n")
    created: list[str] = ["a.txt", "b.txt"]
    try:
        import openpyxl

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "評価"
        sheet["A1"] = "AAA"
        workbook.create_sheet("Sheet2")["D5"] = "BBB"
        workbook.save(os.path.join(corpus, "c.xlsx"))
        created.append("c.xlsx")
    except Exception as exc:
        report.add("xlsx fixture", WARN, f"{type(exc).__name__}")
    try:
        from docx import Document

        document = Document()
        document.add_paragraph("AAA CCC")
        document.save(os.path.join(corpus, "d.docx"))
        created.append("d.docx")
    except Exception as exc:
        report.add("docx fixture", WARN, f"{type(exc).__name__}")
    try:
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "AAA"
        presentation.save(os.path.join(corpus, "e.pptx"))
        created.append("e.pptx")
    except Exception as exc:
        report.add("pptx fixture", WARN, f"{type(exc).__name__}")
    report.add("fixture corpus", PASS, f"{len(created)} files in TEMP")
    return corpus


def _search(root: str, query: str, *, database=None) -> tuple[set[str], object]:
    config = SearchConfig(
        roots=(root,),
        query=query,
        mode="standard" if database else "full",
        workers=2,
        index_enabled=database is not None,
        pdf_ocr_mode="off",
    )
    session = SearchSession(config, database=database)

    def drain() -> None:
        while not session.finished or not session.events.empty():
            try:
                session.events.get(timeout=0.05)
            except Exception:
                continue

    thread = threading.Thread(target=drain, daemon=True)
    thread.start()
    session.start()
    if not session.wait(120):
        session.cancel()
    thread.join(2)
    summary = session.summary
    matched = {os.path.basename(result.path) for result in summary.results} if summary else set()
    return matched, summary


def _check_search(report: Report, corpus: str, workspace: str) -> None:
    matched, summary = _search(corpus, "AAA")
    expected = {"a.txt", "c.xlsx", "d.docx", "e.pptx"}
    ok = expected.issubset(matched)
    report.add(
        "basic search",
        PASS if ok else FAIL,
        f"matched={sorted(matched)}",
    )
    and_matched, _ = _search(corpus, "AAA&BBB")
    report.add(
        "file-level AND (xlsx sheets)",
        PASS if "c.xlsx" in and_matched else FAIL,
        f"matched={sorted(and_matched)}",
    )
    short, _ = _search(corpus, "電源")
    report.add("japanese 2-char search", PASS if "a.txt" in short else FAIL, f"matched={sorted(short)}")
    empty, _ = _search(corpus, "存在しない語句XYZ")
    report.add("zero-hit search", PASS if not empty else FAIL, "0 hits as expected")
    coverage = getattr(summary, "coverage", None)
    if coverage is not None:
        report.add(
            "coverage counters",
            PASS if coverage.discovered >= 3 else WARN,
            coverage.line(),
        )


def _check_index_equivalence(report: Report, corpus: str, workspace: str) -> None:
    path = os.path.join(workspace, "index.sqlite3")
    database = IndexDatabase(
        path,
        max_bytes=64 * 1024 * 1024,
        extractor_version=extraction_fingerprint(
            search_formula=True, include_archives=False, ocr_mode="off", ocr_languages="jpn+eng"
        ),
    )
    try:
        direct, _ = _search(corpus, "AAA&BBB")
        _search(corpus, "AAA", database=database)          # build
        indexed, _ = _search(corpus, "AAA&BBB", database=database)
        report.add(
            "index build",
            PASS if database.status().files >= 3 else FAIL,
            f"{database.status().files} files / {database.status().chunks} chunks",
        )
        report.add(
            "index/direct equivalence",
            PASS if direct == indexed else FAIL,
            f"direct={sorted(direct)} indexed={sorted(indexed)}",
        )
    finally:
        database.close()


def _check_staging(report: Report, workspace: str) -> None:
    from .platform import tempfiles

    source = os.path.join(workspace, "stage-source.txt")
    with open(source, "w", encoding="utf-8") as handle:
        handle.write("staging")
    manager = tempfiles.TempManager(reserve_bytes=0, max_stage_bytes=1024 * 1024)
    try:
        with manager.staged_copy(source, size=os.path.getsize(source), suffix=".txt") as staged:
            staged_path = staged
            with open(staged, "rb") as handle:
                _ = handle.read()
        cleaned = not os.path.exists(staged_path)
        report.add("temp staging cleanup", PASS if cleaned else FAIL, os.path.dirname(staged_path))
    except Exception as exc:
        report.add("temp staging cleanup", FAIL, f"{type(exc).__name__}: {exc}")


def _check_ocr(report: Report, workspace: str) -> None:
    status = tesseract_module.probe()
    if not status.ready:
        report.add("tesseract discovery", SKIP, status.message)
        report.add("jpn+eng languages", SKIP, "Tesseract not installed")
        report.add("ocr scan pdf", SKIP, "Tesseract not installed")
        return
    languages = ", ".join(status.languages)
    report.add("tesseract discovery", PASS, status.executable)
    report.add(
        "jpn+eng languages",
        PASS if {"jpn", "eng"} <= set(status.languages) else WARN,
        languages,
    )
    try:
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (1000, 300), "white")
        ImageDraw.Draw(image).text((50, 120), "AAA 1234", fill="black")
        scan = os.path.join(workspace, "scan.pdf")
        image.save(scan, "PDF", resolution=150.0)
        import pytesseract

        from .extractors import pdf as pdf_module

        rendered = pdf_module.render_page(scan, 0, 2.0)
        text = pytesseract.image_to_string(rendered, lang=status.language_expression or "eng", timeout=30)
        report.add(
            "ocr scan pdf",
            PASS if "AAA" in text.upper().replace(" ", "") else WARN,
            text.strip().replace("\n", " ")[:60],
        )
    except Exception as exc:
        report.add("ocr scan pdf", WARN, f"{type(exc).__name__}: {exc}")


def main(argv: list[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(description="FileScope release self test")
    parser.add_argument("--keep-workspace", action="store_true", help="do not delete the TEMP workspace")
    parser.add_argument("--quiet", action="store_true", help="only print the summary line")
    parser.add_argument("--out", help="also write the report to this file")
    args = parser.parse_args(argv)
    started = time.time()
    report = run(keep_workspace=args.keep_workspace)
    text = report.as_text()
    if args.quiet:
        text = "\n".join(line for line in text.splitlines() if line.startswith("summary:"))
    print(text)
    elapsed_line = f"elapsed: {time.time() - started:.2f}s"
    print(elapsed_line)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as handle:
            handle.write(text + "\n" + elapsed_line + "\n")
    return 1 if report.failed else 0


_ = paths
