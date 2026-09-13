from __future__ import annotations

import csv
import ctypes
import datetime as dt
import io
import itertools
import json
import os
import queue
import re
import shutil
import subprocess
import tempfile
import threading
import time
import traceback
import unicodedata
from collections import defaultdict
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# -----------------------------------------------------------------------------
# Optional document readers. The program itself never sends file contents or
# search terms to the network. Network access only happens indirectly when the
# selected path itself is a network drive / SharePoint-OneDrive synced folder.
# -----------------------------------------------------------------------------
try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    load_workbook = None
    get_column_letter = None

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    from pyxlsb import open_workbook as open_xlsb_workbook
except ImportError:
    open_xlsb_workbook = None

try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    import win32com.client as win32
except ImportError:
    win32 = None


APP_NAME = "FileScope"
EXCEL_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".xlsb"}
PDF_EXTS = {".pdf"}
WORD_EXTS = {".docx", ".docm"}
POWERPOINT_EXTS = {".pptx", ".pptm"}
TEXT_EXTS = {
    ".txt", ".csv", ".tsv", ".log", ".md", ".ini", ".json", ".xml",
    ".html", ".htm", ".py", ".bat", ".ps1", ".sql", ".yaml", ".yml",
    ".cfg", ".conf", ".properties", ".rtf",
}
ALL_SUPPORTED_EXTS = EXCEL_EXTS | PDF_EXTS | WORD_EXTS | POWERPOINT_EXTS | TEXT_EXTS
TEXT_PROBE_BYTES = 64 * 1024
UNKNOWN_TEXT_MAX_BYTES = 50 * 1024 * 1024
OBVIOUS_BINARY_EXTS = {
    ".exe", ".dll", ".bin", ".iso", ".zip", ".7z", ".rar", ".tar", ".gz", ".png", ".jpg", ".jpeg",
    ".gif", ".bmp", ".tif", ".tiff", ".webp", ".mp3", ".wav", ".mp4", ".mov", ".avi", ".db", ".sqlite",
}
HYPHEN_CHARS = "-‐‑‒–—―−ーｰ－"

DEFAULT_WORKERS = 4
INITIAL_DISPLAY_ROWS = 500
LOAD_MORE_ROWS = 500
KEYWORD_DETAIL_ROWS = 1000
UI_REFRESH_MS = 150
AUTOSAVE_FLUSH_EVERY = 100
OCR_PAGE_TIMEOUT_SECONDS = 30

# Safety limits for PCs with relatively little free space.
# Remote Office/PDF files larger than this are parsed in-place rather than
# duplicated into the system TEMP directory. With four workers this caps the
# tool-created staging footprint to about 1 GiB in the usual worst case.
MAX_REMOTE_STAGE_BYTES = 256 * 1024 * 1024
TEMP_FREE_RESERVE_BYTES = 4 * 1024 * 1024 * 1024
AUTOSAVE_LOW_DISK_STOP_BYTES = 3 * 1024 * 1024 * 1024

# Keep the UI/RAM bounded even for pathological queries (for example a common
# word that matches almost every line/cell). Every hit is still streamed to the
# autosave CSV; only the in-memory interactive result set is capped.
MAX_IN_MEMORY_HITS = 250_000


class SearchStopped(Exception):
    pass


@dataclass(frozen=True)
class SearchOptions:
    roots: tuple[str, ...]
    query: str
    exclude_query: str
    include_subfolders: bool
    include_excel: bool
    include_pdf: bool
    include_word: bool
    include_powerpoint: bool
    include_text: bool
    include_unknown_text: bool
    pdf_ocr_mode: str
    case_sensitive: bool
    ignore_width: bool
    part_number_mode: bool
    search_formula: bool
    search_mode: str
    search_path_names: bool
    stage_remote_files: bool
    workers: int = DEFAULT_WORKERS


@dataclass
class SearchHit:
    file_type: str
    path: str
    place: str
    cell: str
    value: str
    matched_keywords: tuple[str, ...] = field(default_factory=tuple)
    severity: str = "HIT"  # HIT / INFO / ERROR


@dataclass
class RuntimeStats:
    total_files: int = 0
    completed_files: int = 0
    hit_count: int = 0
    hit_files: int = 0
    errors: int = 0
    started_at: float = 0.0
    finished_at: float = 0.0
    phase: str = "待機中"
    last_file: str = ""
    active_files: dict[str, tuple[str, float]] = field(default_factory=dict)
    file_type_seconds: dict[str, float] = field(default_factory=lambda: defaultdict(float))


# -----------------------------------------------------------------------------
# Search matcher
# -----------------------------------------------------------------------------
class SearchMatcher:
    """Immutable matcher for one search run.

    Part-number mode intentionally treats hyphen variants and whitespace as
    equivalent, and supports '*' as a wildcard. Normal substring matching is
    always tried first so the special mode never makes ordinary searches worse.
    """

    def __init__(
        self,
        query: str,
        exclude_query: str,
        *,
        case_sensitive: bool,
        ignore_width: bool,
        part_number_mode: bool,
        search_mode: str,
    ):
        self.query = query.strip()
        self.exclude_query = exclude_query.strip()
        self.case_sensitive = case_sensitive
        self.ignore_width = ignore_width
        self.part_number_mode = part_number_mode
        self.search_mode = "AND" if str(search_mode).upper() == "AND" else "OR"

        # Threshold syntax for "N of these terms" searches, e.g.
        #   2of(AAA,落下,耐久)
        # This is intentionally file-level: the matching terms may live in
        # different cells, sheets, pages, slides, or lines of the same file.
        self.threshold_min: int | None = None
        self.threshold_terms_raw: list[str] = []
        threshold = re.fullmatch(r"\s*(\d+)\s*of\s*\((.*)\)\s*", self.query, flags=re.IGNORECASE)
        if threshold:
            raw_terms = [t.strip() for t in re.split(r"[,、;；]+", threshold.group(2)) if t.strip()]
            requested = int(threshold.group(1))
            if raw_terms and 1 <= requested <= len(raw_terms):
                self.threshold_min = requested
                self.threshold_terms_raw = raw_terms
                self.explicit_ops = True
                self.groups_raw = [[term] for term in raw_terms]
            else:
                self.explicit_ops = any(op in self.query for op in [",", "、", "&", ";", "＆"])
                self.groups_raw = self._parse_query(self.query)
        else:
            self.explicit_ops = any(op in self.query for op in [",", "、", "&", ";", "＆"])
            self.groups_raw = self._parse_query(self.query)

        if self.threshold_min is None and not self.explicit_ops and self.groups_raw:
            flat = [term for group in self.groups_raw for term in group]
            if self.search_mode == "OR":
                self.groups_raw = [[term] for term in flat]
            else:
                self.groups_raw = [flat]

        self.groups = [
            [self._prepare_term(term) for term in group]
            for group in self.groups_raw
            if group
        ]
        self.exclude_terms = [
            self._prepare_term(term)
            for term in self._split_plain_terms(self.exclude_query)
            if term
        ]
        self.display_names = [
            group[0] if len(group) == 1 else "&".join(group)
            for group in self.groups_raw
        ]
        self.threshold_display = (
            f"{self.threshold_min}of({','.join(self.threshold_terms_raw)})"
            if self.threshold_min is not None else ""
        )

    def _normalize(self, text: object) -> str:
        if text is None:
            return ""
        value = str(text)
        if self.ignore_width:
            value = unicodedata.normalize("NFKC", value)
        if not self.case_sensitive:
            value = value.casefold()
        return value

    def _canonical_part(self, text: object, *, keep_star: bool = False) -> str:
        if text is None:
            return ""
        value = unicodedata.normalize("NFKC", str(text))
        if not self.case_sensitive:
            value = value.casefold()

        out: list[str] = []
        for ch in value:
            if ch in HYPHEN_CHARS or ch.isspace():
                continue
            if keep_star and ch == "*":
                out.append(ch)
                continue
            if ch.isalnum():
                out.append(ch)
        return "".join(out)

    def _prepare_term(self, raw: str) -> tuple[str, str, re.Pattern[str] | None]:
        norm = self._normalize(raw)
        part = self._canonical_part(raw, keep_star=True) if self.part_number_mode else ""
        wildcard_re = None
        if self.part_number_mode and "*" in part:
            wildcard_re = re.compile(re.escape(part).replace(r"\*", ".*"))
        return norm, part, wildcard_re

    @staticmethod
    def _split_plain_terms(query: str) -> list[str]:
        query = query.strip().replace("　", " ")
        if not query:
            return []
        return [term for term in re.split(r"\s+", query) if term]

    def _parse_query(self, query: str) -> list[list[str]]:
        query = query.strip().replace("　", " ")
        if not query:
            return []

        replacements = {
            "＆": "&", "，": ",", "、": ",", "；": ";", "｜": "|", "／": "/",
        }
        for old, new in replacements.items():
            query = query.replace(old, new)

        query = re.sub(r"\b(かつ|且つ|and)\b", "&", query, flags=re.IGNORECASE)
        query = re.sub(r"\b(または|又は|もしくは|or)\b", ",", query, flags=re.IGNORECASE)

        if any(op in query for op in [",", ";", "&"]):
            groups: list[list[str]] = []
            for clause in re.split(r"[;,]+", query):
                clause = clause.strip()
                if not clause:
                    continue
                if "&" in clause:
                    terms: list[str] = []
                    for part in clause.split("&"):
                        terms.extend(self._split_plain_terms(part))
                else:
                    # Historical behavior: spaces inside an explicit OR clause are AND.
                    terms = self._split_plain_terms(clause)
                if terms:
                    groups.append(terms)
            return groups

        terms = self._split_plain_terms(query)
        return [terms] if terms else []

    def _make_targets(self, candidates: list[object] | tuple[object, ...]) -> tuple[str, str]:
        values = [c for c in candidates if c is not None and str(c) != ""]
        target_norm = "\n".join(self._normalize(c) for c in values)
        target_part = ""
        if self.part_number_mode:
            target_part = "\n".join(self._canonical_part(c) for c in values)
        return target_norm, target_part

    @staticmethod
    def _term_matches(prepared: tuple[str, str, re.Pattern[str] | None], target_norm: str, target_part: str) -> bool:
        norm, part, wildcard_re = prepared
        if norm and norm in target_norm:
            return True
        if not part:
            return False
        if wildcard_re is not None:
            return wildcard_re.search(target_part) is not None
        return part in target_part

    def evaluate(self, candidates: list[object] | tuple[object, ...]) -> tuple[str, ...]:
        if not self.groups:
            return ()
        target_norm, target_part = self._make_targets(candidates)
        if not target_norm and not target_part:
            return ()

        for term in self.exclude_terms:
            if self._term_matches(term, target_norm, target_part):
                return ()

        if self.threshold_min is not None:
            count = sum(
                1 for group in self.groups
                if group and self._term_matches(group[0], target_norm, target_part)
            )
            return (self.threshold_display,) if count >= self.threshold_min else ()

        matched: list[str] = []
        for display, group in zip(self.display_names, self.groups):
            if all(self._term_matches(term, target_norm, target_part) for term in group):
                matched.append(display)
        return tuple(matched)

    def unit_terms(self, candidates: list[object] | tuple[object, ...]) -> tuple[str, ...]:
        """Return individual search terms that match this unit. Used for file-level AND."""
        if not self.groups:
            return ()
        target_norm, target_part = self._make_targets(candidates)
        if not target_norm and not target_part:
            return ()
        for term in self.exclude_terms:
            if self._term_matches(term, target_norm, target_part):
                return ()
        matched: list[str] = []
        seen: set[str] = set()
        for group_raw, group in zip(self.groups_raw, self.groups):
            for raw, prepared in zip(group_raw, group):
                if raw not in seen and self._term_matches(prepared, target_norm, target_part):
                    matched.append(raw)
                    seen.add(raw)
        return tuple(matched)

    def file_displays(self, present_terms: set[str]) -> tuple[str, ...]:
        """Return query expressions satisfied somewhere in the same file."""
        if self.threshold_min is not None:
            present = sum(1 for raw in self.threshold_terms_raw if raw in present_terms)
            return (self.threshold_display,) if present >= self.threshold_min else ()
        matched: list[str] = []
        for display, group_raw in zip(self.display_names, self.groups_raw):
            if all(raw in present_terms for raw in group_raw):
                matched.append(display)
        return tuple(matched)

    def terms_for_displays(self, displays: tuple[str, ...]) -> set[str]:
        if self.threshold_min is not None and self.threshold_display in set(displays):
            return set(self.threshold_terms_raw)
        wanted: set[str] = set()
        display_set = set(displays)
        for display, group_raw in zip(self.display_names, self.groups_raw):
            if display in display_set:
                wanted.update(group_raw)
        return wanted

    def contains_excluded(self, candidates: list[object] | tuple[object, ...]) -> bool:
        if not self.exclude_terms:
            return False
        target_norm, target_part = self._make_targets(candidates)
        return any(self._term_matches(term, target_norm, target_part) for term in self.exclude_terms)

    def snippet(self, text: object, max_len: int = 260) -> str:
        value = " ".join(str(text or "").replace("\r", "\n").split())
        if len(value) <= max_len:
            return value

        norm = self._normalize(value)
        positions: list[int] = []
        for group in self.groups:
            for term, _, _ in group:
                if term:
                    pos = norm.find(term)
                    if pos >= 0:
                        positions.append(pos)
        pos = min(positions) if positions else 0
        half = max_len // 2
        start = max(0, pos - half)
        end = min(len(value), start + max_len)
        prefix = "..." if start else ""
        suffix = "..." if end < len(value) else ""
        return prefix + value[start:end] + suffix


# -----------------------------------------------------------------------------
# File helpers
# -----------------------------------------------------------------------------
def get_config_path() -> str:
    if os.name == "nt":
        base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    else:
        base = os.path.join(os.path.expanduser("~"), ".config")
    folder = os.path.join(base, APP_NAME)
    try:
        os.makedirs(folder, exist_ok=True)
        return os.path.join(folder, "settings.json")
    except Exception:
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), "search_tool_settings_v2.json")


CONFIG_PATH = get_config_path()
LEGACY_CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "search_tool_settings.json")


def is_remote_path(path: str) -> bool:
    """Return True for UNC or Windows mapped network drives."""
    try:
        path = os.path.abspath(path)
    except Exception:
        pass
    if str(path).startswith("\\\\"):
        return True
    if os.name != "nt":
        return False
    try:
        drive, _ = os.path.splitdrive(path)
        if not drive:
            return False
        DRIVE_REMOTE = 4
        return ctypes.windll.kernel32.GetDriveTypeW(drive + "\\") == DRIVE_REMOTE
    except Exception:
        return False


def find_tesseract_executable() -> str | None:
    """Find Tesseract without requiring users to edit PATH manually."""
    if pytesseract is None:
        return None
    candidates: list[str] = []
    env_cmd = os.environ.get("TESSERACT_CMD", "").strip()
    if env_cmd:
        candidates.append(env_cmd)
    which = shutil.which("tesseract")
    if which:
        candidates.append(which)
    app_dir = os.path.dirname(os.path.abspath(__file__))
    candidates.extend([
        os.path.join(app_dir, "Tesseract-OCR", "tesseract.exe"),
        os.path.join(app_dir, "tesseract", "tesseract.exe"),
        os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"), "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"), "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Tesseract-OCR", "tesseract.exe"),
    ])
    for candidate in candidates:
        if candidate and os.path.isfile(candidate):
            try:
                pytesseract.pytesseract.tesseract_cmd = candidate
                return candidate
            except Exception:
                pass
    return None


def detect_ocr_capability() -> tuple[bool, str, str]:
    """Return (ready, language_expression, human_status)."""
    if pdfium is None:
        return False, "", "OCR未導入: pypdfium2 が必要"
    if pytesseract is None:
        return False, "", "OCR未導入: pytesseract が必要"
    cmd = find_tesseract_executable()
    if not cmd:
        return False, "", "OCR未導入: Tesseract本体が見つかりません"
    try:
        languages = set(pytesseract.get_languages(config=""))
    except Exception as exc:
        return False, "", f"OCR確認エラー: {type(exc).__name__}"
    if "jpn" in languages and "eng" in languages:
        return True, "jpn+eng", "OCR利用可: 日本語+英語"
    if "jpn" in languages:
        return True, "jpn", "OCR利用可: 日本語（eng未導入）"
    if "eng" in languages:
        return True, "eng", "OCR利用可: 英語のみ（jpn未導入）"
    return False, "", "OCR未導入: jpn/eng 言語データがありません"


@contextmanager
def local_parse_path(path: str, options: SearchOptions):
    """For remote structured documents, copy once then parse locally.

    XLSX/DOCX/PPTX are ZIP containers. Parsing them directly on SMB can cause
    many small network reads. A sequential copy to a local temp file is usually
    more stable and often much faster. Original files are never modified.
    """
    ext = os.path.splitext(path)[1].lower()
    should_stage = options.stage_remote_files and is_remote_path(path) and ext not in TEXT_EXTS
    if not should_stage:
        yield path
        return

    # Protect small system drives. A very large online/SMB file is safer to
    # parse directly than to create another full copy on C:.
    try:
        source_size = os.path.getsize(path)
    except Exception:
        source_size = 0

    if source_size > MAX_REMOTE_STAGE_BYTES:
        yield path
        return

    try:
        temp_dir = tempfile.gettempdir()
        free_bytes = shutil.disk_usage(temp_dir).free
        required = max(source_size, 1) + TEMP_FREE_RESERVE_BYTES
        if free_bytes < required:
            yield path
            return
    except Exception:
        pass

    temp_path = ""
    try:
        fd, temp_path = tempfile.mkstemp(prefix="file_search_v2_", suffix=ext)
        os.close(fd)
        shutil.copyfile(path, temp_path)
        yield temp_path
    finally:
        if temp_path:
            try:
                os.remove(temp_path)
            except Exception:
                pass


def excel_cell_candidates(value: object, number_format: str = "") -> list[str]:
    if value is None:
        return []

    candidates = [str(value)]

    if isinstance(value, bool):
        candidates.extend(["TRUE" if value else "FALSE", "true" if value else "false"])
    elif isinstance(value, int):
        candidates.extend([
            f"{value}.0", f"{value}.00", f"{value:03d}", f"{value:04d}", f"{value:05d}",
        ])
    elif isinstance(value, float):
        if value.is_integer():
            iv = int(value)
            candidates.extend([
                str(iv), f"{iv}.0", f"{iv}.00", f"{iv:03d}", f"{iv:04d}", f"{iv:05d}",
            ])
        decimals = decimal_places_from_format(number_format)
        if decimals is not None:
            candidates.append(f"{value:.{decimals}f}")
    elif isinstance(value, (dt.datetime, dt.date)):
        candidates.extend([
            value.strftime("%Y/%m/%d"), value.strftime("%Y-%m-%d"),
            f"{value.year}/{value.month}/{value.day}",
        ])
        if isinstance(value, dt.datetime):
            candidates.extend([
                value.strftime("%Y/%m/%d %H:%M"), value.strftime("%Y/%m/%d %H:%M:%S"),
            ])

    return list(dict.fromkeys(candidates))


def decimal_places_from_format(fmt: str) -> int | None:
    if not fmt or "." not in str(fmt):
        return None
    after = str(fmt).split(".", 1)[1]
    count = 0
    for ch in after:
        if ch in ("0", "#"):
            count += 1
        else:
            break
    return min(count, 10) if count else None



def looks_like_text_file(path: str) -> bool:
    """Lightweight local-only probe for proprietary text-like files (.td2/.dat/.prm etc.)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in OBVIOUS_BINARY_EXTS:
        return False
    try:
        size = os.path.getsize(path)
        if size > UNKNOWN_TEXT_MAX_BYTES:
            return False
        with open(path, "rb") as f:
            sample = f.read(TEXT_PROBE_BYTES)
    except Exception:
        return False
    if not sample:
        return True
    nul_ratio = sample.count(b"\x00") / max(len(sample), 1)
    if nul_ratio > 0.01:
        return False
    for enc in ("utf-8-sig", "utf-8", "cp932", "shift_jis", "utf-16", "utf-16le", "utf-16be"):
        try:
            text = sample.decode(enc)
            break
        except Exception:
            text = ""
    if not text:
        text = sample.decode("utf-8", errors="ignore")
    if not text:
        return False
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\r\n\t")
    return (printable / max(len(text), 1)) >= 0.75


@dataclass
class TextEvidence:
    file_type: str
    path: str
    place: str
    cell: str
    value: str
    matched_terms: tuple[str, ...]


def detect_text_encoding(raw: io.BufferedReader) -> str:
    sample = raw.read(65536)
    raw.seek(0)
    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    try:
        sample.decode("utf-8", errors="strict")
        return "utf-8"
    except UnicodeDecodeError:
        pass

    # UTF-16 files without BOM often contain many NUL bytes.
    if sample and sample.count(b"\x00") / len(sample) > 0.15:
        for enc in ("utf-16le", "utf-16be"):
            try:
                sample.decode(enc, errors="strict")
                return enc
            except UnicodeDecodeError:
                pass
    return "cp932"


# -----------------------------------------------------------------------------
# Main application
# -----------------------------------------------------------------------------
class SearchApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("FileScope - Excel / PDF / Office / Text 横断検索")
        self.root.geometry("1280x760")
        self.root.minsize(820, 520)
        self.root.resizable(True, True)

        # UI variables
        self.folder_var = tk.StringVar()
        self.keyword_var = tk.StringVar()
        self.exclude_keyword_var = tk.StringVar()
        self.include_subfolders_var = tk.BooleanVar(value=True)
        self.include_excel_var = tk.BooleanVar(value=True)
        self.include_pdf_var = tk.BooleanVar(value=True)
        self.include_word_var = tk.BooleanVar(value=True)
        self.include_powerpoint_var = tk.BooleanVar(value=True)
        self.include_text_var = tk.BooleanVar(value=True)
        self.include_unknown_text_var = tk.BooleanVar(value=True)
        self.pdf_ocr_mode_var = tk.StringVar(value="自動")
        self.case_sensitive_var = tk.BooleanVar(value=False)
        self.ignore_width_var = tk.BooleanVar(value=True)
        self.part_number_mode_var = tk.BooleanVar(value=False)
        self.search_formula_var = tk.BooleanVar(value=True)
        self.search_mode_var = tk.StringVar(value="OR")
        self.search_path_names_var = tk.BooleanVar(value=True)
        self.stage_remote_files_var = tk.BooleanVar(value=True)
        self.status_var = tk.StringVar(value="待機中")
        self.progress_var = tk.IntVar(value=0)
        self.summary_var = tk.StringVar(value="0 / 0 ファイル   0 ヒット")

        self.history: list[str] = []
        self.results: list[SearchHit] = []
        self.keyword_results: dict[str, list[int]] = defaultdict(list)
        self.file_hit_counts: dict[str, int] = defaultdict(int)
        self.type_hit_counts: dict[str, int] = defaultdict(int)
        self.keyword_hit_counts: dict[str, int] = defaultdict(int)
        self.hit_files: set[str] = set()
        self.confirmed_files: set[str] = set()
        self._sort_column = ""
        self._sort_desc = False
        self._keyword_sort_column = ""
        self._keyword_sort_desc = False
        self.results_lock = threading.RLock()
        self.stats = RuntimeStats()
        self.stats_lock = threading.Lock()
        self.stop_event = threading.Event()
        self.pause_event = threading.Event()
        self.pause_event.set()
        self._search_running = False
        self._search_id = 0
        self._search_thread: threading.Thread | None = None
        # pypdf can consume a lot of memory on pathological pages. Keep PDF
        # extraction single-file-at-a-time while other file types still run in parallel.
        self._pdf_semaphore = threading.Semaphore(1)

        self._display_limit = INITIAL_DISPLAY_ROWS
        self._rendered_result_count = 0
        self._result_tree_map: dict[str, int] = {}
        self._keyword_tree_map: dict[str, int] = {}
        self._keyword_list_keys: list[str] = []
        self._last_keyword_refresh = 0.0
        self._ui_job = None
        self._active_options: SearchOptions | None = None
        self._matcher: SearchMatcher | None = None

        self._autosave_path: str | None = None
        self._autosave_fp = None
        self._autosave_writer = None
        self._autosave_count = 0
        self._autosave_lock = threading.Lock()
        self._memory_hit_limit_reached = False
        self._low_disk_stop = False
        self._ocr_ready, self._ocr_lang, self._ocr_status = detect_ocr_capability()

        self.load_settings()
        self.apply_theme()
        self.build_ui()
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        self._schedule_ui_refresh()

    # ------------------------------------------------------------------ UI
    def apply_theme(self):
        style = ttk.Style()
        for name in ("vista", "winnative", "clam", "alt", "default"):
            if name in style.theme_names():
                try:
                    style.theme_use(name)
                    break
                except Exception:
                    pass
        style.configure("Treeview", rowheight=26)
        style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"))
        style.configure("TButton", padding=4)

    def build_ui(self):
        main = ttk.Frame(self.root, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        folder_row = ttk.Frame(main)
        folder_row.pack(fill=tk.X)
        ttk.Label(folder_row, text="検索フォルダ:").pack(side=tk.LEFT)
        ttk.Entry(folder_row, textvariable=self.folder_var).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=6)
        ttk.Button(folder_row, text="参照", command=self.select_folder).pack(side=tk.LEFT)
        ttk.Button(folder_row, text="追加", command=self.append_folder).pack(side=tk.LEFT, padx=(5, 0))

        query_row = ttk.Frame(main)
        query_row.pack(fill=tk.X, pady=(8, 4))
        ttk.Label(query_row, text="検索文字:").pack(side=tk.LEFT)
        self.keyword_combo = ttk.Combobox(query_row, textvariable=self.keyword_var, values=self.history)
        self.keyword_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=6)
        ttk.Label(query_row, text="除外:").pack(side=tk.LEFT, padx=(8, 0))
        ttk.Entry(query_row, textvariable=self.exclude_keyword_var, width=30).pack(side=tk.LEFT, padx=6)
        ttk.Label(query_row, text="空白区切り:").pack(side=tk.LEFT, padx=(8, 2))
        ttk.Combobox(query_row, textvariable=self.search_mode_var, values=("OR", "AND"), width=5, state="readonly").pack(side=tk.LEFT)

        ttk.Label(
            main,
            text='検索式: A B＝右のOR/AND、A&B＝AND、A,B＝OR、2of(A,B,C)＝3語中2語以上。除外欄＝除外。部品番号はハイフン/空白差を吸収、*＝ワイルドカード。',
        ).pack(fill=tk.X, pady=(0, 5))

        type_frame = ttk.LabelFrame(main, text="検索対象")
        type_frame.pack(fill=tk.X, pady=3)
        for text, var in [
            ("Excel", self.include_excel_var), ("PDF", self.include_pdf_var),
            ("Word", self.include_word_var), ("PowerPoint", self.include_powerpoint_var),
            ("Text/CSV/Log", self.include_text_var),
            ("未知Text", self.include_unknown_text_var),
        ]:
            ttk.Checkbutton(type_frame, text=text, variable=var).pack(side=tk.LEFT, padx=7, pady=3)
        ttk.Checkbutton(type_frame, text="子フォルダ", variable=self.include_subfolders_var).pack(side=tk.LEFT, padx=7)
        ttk.Label(type_frame, text="PDF OCR:").pack(side=tk.LEFT, padx=(12, 2))
        ttk.Combobox(type_frame, textvariable=self.pdf_ocr_mode_var, values=("OFF", "自動", "全ページ"), width=7, state="readonly").pack(side=tk.LEFT)
        self.ocr_status_label = ttk.Label(type_frame, text=self._ocr_status)
        self.ocr_status_label.pack(side=tk.LEFT, padx=(8, 0))

        option_frame = ttk.LabelFrame(main, text="検索オプション")
        option_frame.pack(fill=tk.X, pady=3)
        for text, var in [
            ("半角/全角を区別しない", self.ignore_width_var),
            ("大文字/小文字を区別", self.case_sensitive_var),
            ("部品番号モード", self.part_number_mode_var),
            ("Excel数式も検索", self.search_formula_var),
            ("ファイル名/フォルダ名も検索", self.search_path_names_var),
            ("共有ドライブは一時コピーして解析", self.stage_remote_files_var),
        ]:
            ttk.Checkbutton(option_frame, text=text, variable=var).pack(side=tk.LEFT, padx=6, pady=3)

        action_row = ttk.Frame(main)
        action_row.pack(fill=tk.X, pady=(6, 4))
        self.search_button = ttk.Button(action_row, text="検索開始", command=self.start_search)
        self.search_button.pack(side=tk.LEFT, padx=(0, 5))
        self.pause_button = ttk.Button(action_row, text="一時停止", command=self.pause_resume_search, state=tk.DISABLED)
        self.pause_button.pack(side=tk.LEFT, padx=5)
        self.stop_button = ttk.Button(action_row, text="中断", command=self.stop_search, state=tk.DISABLED)
        self.stop_button.pack(side=tk.LEFT, padx=5)
        ttk.Button(action_row, text="結果CSV保存", command=self.save_csv).pack(side=tk.LEFT, padx=(14, 5))
        ttk.Button(action_row, text="設定保存", command=self.save_settings_button).pack(side=tk.LEFT, padx=5)
        ttk.Label(action_row, text="外部送信なし / ローカル処理", foreground="#555555").pack(side=tk.RIGHT)

        progress_row = ttk.Frame(main)
        progress_row.pack(fill=tk.X, pady=(0, 6))
        self.progress_bar = ttk.Progressbar(progress_row, variable=self.progress_var, mode="determinate", maximum=1)
        self.progress_bar.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Label(progress_row, textvariable=self.summary_var, width=38, anchor=tk.E).pack(side=tk.RIGHT, padx=(8, 0))
        ttk.Label(main, textvariable=self.status_var).pack(fill=tk.X, pady=(0, 5))

        self.notebook = ttk.Notebook(main)
        self.notebook.pack(fill=tk.BOTH, expand=True)
        self.build_results_tab()
        self.build_keyword_tab()

        self.context_menu = tk.Menu(self.root, tearoff=0)
        self.context_menu.add_command(label="ファイルを開く", command=self.open_selected_file)
        self.context_menu.add_command(label="フォルダを開く", command=self.open_selected_folder)
        self.context_menu.add_command(label="Excelの該当セルを開く", command=self.open_selected_excel_cell)
        self.context_menu.add_separator()
        self.context_menu.add_command(label="内容をコピー", command=self.copy_selected_value)
        self.context_menu.add_command(label="ファイルパスをコピー", command=self.copy_selected_path)
        self.context_menu.add_separator()
        self.context_menu.add_command(label="確認済み状態を切り替え", command=self.toggle_confirmed_state)

    def build_results_tab(self):
        tab = ttk.Frame(self.notebook)
        self.notebook.add(tab, text="結果一覧")

        columns = ("confirmed", "file", "keywords", "place", "cell", "value", "type", "path")
        self.result_tree = ttk.Treeview(tab, columns=columns, show="headings", selectmode="browse")
        specs = {
            "confirmed": ("済", 42), "file": ("ファイル", 220), "keywords": ("マッチ", 135), "place": ("場所", 135),
            "cell": ("セル", 65), "value": ("内容", 380), "type": ("種類", 80), "path": ("パス", 230),
        }
        for col, (label, width) in specs.items():
            self.result_tree.heading(col, text=label, command=lambda c=col: self.sort_results_by(c))
            self.result_tree.column(col, width=width, minwidth=50, stretch=(col in {"file", "value", "path"}))
        self.result_tree.pack(side=tk.TOP, fill=tk.BOTH, expand=True)
        self.result_tree.bind("<Button-3>", self.show_context_menu)
        self.result_tree.bind("<Double-1>", lambda _e: self.open_selected_smart())

        bottom = ttk.Frame(tab)
        bottom.pack(fill=tk.X, pady=(5, 0))
        self.display_info_var = tk.StringVar(value="表示 0 / 0")
        ttk.Label(bottom, textvariable=self.display_info_var).pack(side=tk.LEFT)
        ttk.Button(bottom, text=f"さらに{LOAD_MORE_ROWS}件表示", command=self.load_more_results).pack(side=tk.RIGHT)

    def build_keyword_tab(self):
        tab = ttk.Frame(self.notebook)
        self.notebook.add(tab, text="キーワード別")
        pane = ttk.Panedwindow(tab, orient=tk.HORIZONTAL)
        pane.pack(fill=tk.BOTH, expand=True)

        left = ttk.Frame(pane, padding=5)
        right = ttk.Frame(pane, padding=5)
        pane.add(left, weight=1)
        pane.add(right, weight=4)

        ttk.Label(left, text="検索語 / ヒット数").pack(fill=tk.X)
        self.keyword_listbox = tk.Listbox(left, exportselection=False)
        self.keyword_listbox.pack(fill=tk.BOTH, expand=True, pady=(5, 0))
        self.keyword_listbox.bind("<<ListboxSelect>>", self.on_keyword_select)

        cols = ("confirmed", "file", "place", "cell", "value", "type", "path")
        self.keyword_tree = ttk.Treeview(right, columns=cols, show="headings", selectmode="browse")
        for col, label, width in [
            ("confirmed", "済", 42), ("file", "ファイル", 220), ("place", "場所", 130), ("cell", "セル", 65),
            ("value", "内容", 380), ("type", "種類", 80), ("path", "パス", 220),
        ]:
            self.keyword_tree.heading(col, text=label, command=lambda c=col: self.sort_keyword_results_by(c))
            self.keyword_tree.column(col, width=width, minwidth=50, stretch=(col in {"file", "value", "path"}))
        self.keyword_tree.pack(fill=tk.BOTH, expand=True)
        self.keyword_tree.bind("<Button-3>", self.show_context_menu)
        self.keyword_tree.bind("<Double-1>", lambda _e: self.open_selected_smart())

    # --------------------------------------------------------------- settings
    def load_settings(self):
        source_path = CONFIG_PATH if os.path.exists(CONFIG_PATH) else LEGACY_CONFIG_PATH
        if not os.path.exists(source_path):
            return
        try:
            with open(source_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            self.folder_var.set(data.get("folder", ""))
            self.keyword_var.set(data.get("keyword", ""))
            self.exclude_keyword_var.set(data.get("exclude_keyword", ""))
            self.include_subfolders_var.set(data.get("include_subfolders", True))
            self.include_excel_var.set(data.get("include_excel", True))
            self.include_pdf_var.set(data.get("include_pdf", True))
            self.include_word_var.set(data.get("include_word", True))
            self.include_powerpoint_var.set(data.get("include_powerpoint", True))
            self.include_text_var.set(data.get("include_text", True))
            self.include_unknown_text_var.set(data.get("include_unknown_text", True))
            self.pdf_ocr_mode_var.set(data.get("pdf_ocr_mode", "自動"))
            self.confirmed_files = {
                os.path.normcase(os.path.normpath(str(p)))
                for p in data.get("confirmed_files", [])
                if p
            }
            self.case_sensitive_var.set(data.get("case_sensitive", False))
            self.ignore_width_var.set(data.get("ignore_width", True))
            self.part_number_mode_var.set(data.get("part_number_mode", False))
            self.search_formula_var.set(data.get("search_formula", True))
            self.search_mode_var.set(data.get("search_mode", "OR"))
            self.search_path_names_var.set(data.get("search_path_names", True))
            self.stage_remote_files_var.set(data.get("stage_remote_files", True))
            self.history = data.get("history", [])[:30]
        except Exception:
            pass

    def save_settings(self):
        data = {
            "folder": self.folder_var.get(), "keyword": self.keyword_var.get(),
            "exclude_keyword": self.exclude_keyword_var.get(),
            "include_subfolders": self.include_subfolders_var.get(),
            "include_excel": self.include_excel_var.get(), "include_pdf": self.include_pdf_var.get(),
            "include_word": self.include_word_var.get(), "include_powerpoint": self.include_powerpoint_var.get(),
            "include_text": self.include_text_var.get(), "include_unknown_text": self.include_unknown_text_var.get(),
            "pdf_ocr_mode": self.pdf_ocr_mode_var.get(), "confirmed_files": sorted(self.confirmed_files),
            "case_sensitive": self.case_sensitive_var.get(),
            "ignore_width": self.ignore_width_var.get(), "part_number_mode": self.part_number_mode_var.get(),
            "search_formula": self.search_formula_var.get(), "search_mode": self.search_mode_var.get(),
            "search_path_names": self.search_path_names_var.get(),
            "stage_remote_files": self.stage_remote_files_var.get(), "history": self.history,
        }
        try:
            os.makedirs(os.path.dirname(CONFIG_PATH), exist_ok=True)
            with open(CONFIG_PATH, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def save_settings_button(self):
        self.add_history(self.keyword_var.get())
        self.save_settings()
        messagebox.showinfo("完了", "設定を保存しました。")

    def add_history(self, keyword: str):
        keyword = keyword.strip()
        if not keyword:
            return
        if keyword in self.history:
            self.history.remove(keyword)
        self.history.insert(0, keyword)
        self.history = self.history[:30]
        if hasattr(self, "keyword_combo"):
            self.keyword_combo["values"] = self.history

    # ----------------------------------------------------------- search setup
    def parse_folder_list(self, text: str) -> list[str]:
        return [p.strip() for p in re.split(r"[;\n\r]+", text or "") if p.strip()]

    def select_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.folder_var.set(folder)

    def append_folder(self):
        folder = filedialog.askdirectory()
        if not folder:
            return
        folders = self.parse_folder_list(self.folder_var.get())
        if folder not in folders:
            folders.append(folder)
            self.folder_var.set(";".join(folders))

    def snapshot_options(self) -> SearchOptions:
        return SearchOptions(
            roots=tuple(self.parse_folder_list(self.folder_var.get())),
            query=self.keyword_var.get().strip(),
            exclude_query=self.exclude_keyword_var.get().strip(),
            include_subfolders=self.include_subfolders_var.get(),
            include_excel=self.include_excel_var.get(), include_pdf=self.include_pdf_var.get(),
            include_word=self.include_word_var.get(), include_powerpoint=self.include_powerpoint_var.get(),
            include_text=self.include_text_var.get(), include_unknown_text=self.include_unknown_text_var.get(),
            pdf_ocr_mode=self.pdf_ocr_mode_var.get(), case_sensitive=self.case_sensitive_var.get(),
            ignore_width=self.ignore_width_var.get(), part_number_mode=self.part_number_mode_var.get(),
            search_formula=self.search_formula_var.get(), search_mode=self.search_mode_var.get(),
            search_path_names=self.search_path_names_var.get(), stage_remote_files=self.stage_remote_files_var.get(),
            workers=DEFAULT_WORKERS,
        )

    def start_search(self):
        if self._search_running:
            return
        options = self.snapshot_options()
        if not options.roots:
            messagebox.showwarning("確認", "検索フォルダを指定してください。")
            return
        if not options.query:
            messagebox.showwarning("確認", "検索文字を入力してください。")
            return
        if not any([options.include_excel, options.include_pdf, options.include_word, options.include_powerpoint, options.include_text]):
            messagebox.showwarning("確認", "検索対象のファイル種類を1つ以上選択してください。")
            return

        self._search_id += 1
        search_id = self._search_id
        self._active_options = options
        self._matcher = SearchMatcher(
            options.query, options.exclude_query,
            case_sensitive=options.case_sensitive, ignore_width=options.ignore_width,
            part_number_mode=options.part_number_mode, search_mode=options.search_mode,
        )
        self.add_history(options.query)
        self.save_settings()

        self.stop_event.clear()
        self.pause_event.set()
        self._search_running = True
        self._display_limit = INITIAL_DISPLAY_ROWS
        self._rendered_result_count = 0
        self._last_keyword_refresh = 0.0
        self._memory_hit_limit_reached = False
        self._low_disk_stop = False

        with self.results_lock:
            self.results.clear()
            self.keyword_results = defaultdict(list)
            self.file_hit_counts = defaultdict(int)
            self.type_hit_counts = defaultdict(int)
            self.keyword_hit_counts = defaultdict(int)
            self.hit_files = set()

        with self.stats_lock:
            self.stats = RuntimeStats(started_at=time.perf_counter(), phase="ファイル一覧取得中...")

        self.clear_ui_results()
        self._start_autosave(options.query)
        self.search_button.config(state=tk.DISABLED)
        self.pause_button.config(state=tk.NORMAL, text="一時停止")
        self.stop_button.config(state=tk.NORMAL)
        self.status_var.set("ファイル一覧取得中...")
        self.progress_var.set(0)
        self.progress_bar.config(maximum=1)

        self._search_thread = threading.Thread(
            target=self.search_coordinator, args=(search_id, options, self._matcher), daemon=True,
            name=f"search-coordinator-{search_id}",
        )
        self._search_thread.start()

    def pause_resume_search(self):
        if not self._search_running:
            return
        if self.pause_event.is_set():
            self.pause_event.clear()
            self.pause_button.config(text="再開")
            self.status_var.set("一時停止中...")
        else:
            self.pause_event.set()
            self.pause_button.config(text="一時停止")

    def stop_search(self):
        if not self._search_running:
            return
        self.stop_event.set()
        self.pause_event.set()
        self.status_var.set("中断要求済み。処理中のファイルが戻り次第停止します...")

    def pause_point(self):
        while not self.pause_event.wait(0.2):
            if self.stop_event.is_set():
                raise SearchStopped()
        if self.stop_event.is_set():
            raise SearchStopped()

    # ------------------------------------------------------------- coordinator
    def target_extensions(self, options: SearchOptions) -> set[str]:
        exts: set[str] = set()
        if options.include_excel:
            exts |= EXCEL_EXTS
        if options.include_pdf:
            exts |= PDF_EXTS
        if options.include_word:
            exts |= WORD_EXTS
        if options.include_powerpoint:
            exts |= POWERPOINT_EXTS
        if options.include_text:
            exts |= TEXT_EXTS
        return exts

    def iter_files(self, search_id: int, options: SearchOptions):
        """Yield candidates as they are discovered instead of building one huge list first.

        This matters especially for SharePoint/OneDrive and very large SMB trees: workers can
        start searching immediately while directory enumeration continues. Unknown extensions
        are *not* opened here; their text/binary probe is deferred to a worker so Files On-Demand
        placeholders are not hydrated during the directory-listing phase.
        """
        target_exts = self.target_extensions(options)
        seen: set[str] = set()

        def accept_path(path: str, name: str) -> bool:
            if name.startswith("~$"):
                return False
            ext = os.path.splitext(name)[1].lower()
            if ext in target_exts:
                return True
            if not (options.include_text and options.include_unknown_text):
                return False
            return ext not in ALL_SUPPORTED_EXTS and ext not in OBVIOUS_BINARY_EXTS

        for root in options.roots:
            if self.stop_event.is_set():
                raise SearchStopped()
            root = os.path.normpath(root)
            try:
                if not os.path.isdir(root):
                    self.add_hit(SearchHit("ERROR", root, "", "", "フォルダを開けません", (), "ERROR"), search_id)
                    continue
            except Exception as exc:
                self.add_hit(SearchHit("ERROR", root, "", "", f"フォルダ確認エラー: {exc}", (), "ERROR"), search_id)
                continue

            if options.include_subfolders:
                def onerror(err):
                    self.add_hit(SearchHit("ERROR", getattr(err, "filename", root) or root, "", "", str(err), (), "ERROR"), search_id)

                for current, _dirs, names in os.walk(root, onerror=onerror):
                    self.pause_point()
                    for name in names:
                        path = os.path.join(current, name)
                        if not accept_path(path, name):
                            continue
                        key = os.path.normcase(os.path.abspath(path))
                        if key in seen:
                            continue
                        seen.add(key)
                        yield path
            else:
                try:
                    for entry in os.scandir(root):
                        self.pause_point()
                        if not entry.is_file() or not accept_path(entry.path, entry.name):
                            continue
                        key = os.path.normcase(os.path.abspath(entry.path))
                        if key in seen:
                            continue
                        seen.add(key)
                        yield entry.path
                except Exception as exc:
                    self.add_hit(SearchHit("ERROR", root, "", "", f"一覧取得エラー: {exc}", (), "ERROR"), search_id)

    def search_coordinator(self, search_id: int, options: SearchOptions, matcher: SearchMatcher):
        work_q: queue.Queue[str | None] = queue.Queue(maxsize=max(options.workers * 3, 8))
        workers: list[threading.Thread] = []
        try:
            # Start consumers first. This removes the old long "一覧取得中だけ" phase.
            for n in range(max(1, min(options.workers, 8))):
                t = threading.Thread(
                    target=self.worker_loop,
                    args=(search_id, options, matcher, work_q),
                    daemon=True,
                    name=f"file-worker-{search_id}-{n + 1}",
                )
                t.start()
                workers.append(t)

            with self.stats_lock:
                self.stats.phase = "列挙＋検索中"
                self.stats.total_files = 0

            for path in self.iter_files(search_id, options):
                if self.stop_event.is_set() or search_id != self._search_id:
                    break
                with self.stats_lock:
                    self.stats.total_files += 1
                while True:
                    try:
                        work_q.put(path, timeout=0.2)
                        break
                    except queue.Full:
                        if self.stop_event.is_set() or search_id != self._search_id:
                            break
                if self.stop_event.is_set() or search_id != self._search_id:
                    break

            with self.stats_lock:
                if search_id == self._search_id and not self.stop_event.is_set():
                    self.stats.phase = "検索中（一覧取得完了）"

            for _ in workers:
                while True:
                    try:
                        work_q.put(None, timeout=0.2)
                        break
                    except queue.Full:
                        if search_id != self._search_id:
                            return

            while any(t.is_alive() for t in workers):
                if search_id != self._search_id:
                    return
                for t in workers:
                    t.join(timeout=0.1)

            if search_id == self._search_id:
                self.finish_search(search_id, stopped=self.stop_event.is_set())

        except SearchStopped:
            if search_id == self._search_id:
                self.finish_search(search_id, stopped=True)
        except Exception:
            if search_id == self._search_id:
                self.add_hit(SearchHit("ERROR", "", "", "", traceback.format_exc(), (), "ERROR"), search_id)
                self.finish_search(search_id, stopped=True)

    def worker_loop(self, search_id: int, options: SearchOptions, matcher: SearchMatcher, work_q: queue.Queue):
        thread_name = threading.current_thread().name
        while True:
            item = work_q.get()
            try:
                if item is None:
                    return
                if search_id != self._search_id or self.stop_event.is_set():
                    continue
                self.pause_point()
                with self.stats_lock:
                    self.stats.active_files[thread_name] = (item, time.perf_counter())
                    self.stats.last_file = item

                started = time.perf_counter()
                try:
                    self.process_file(item, search_id, options, matcher)
                except SearchStopped:
                    pass
                except Exception as exc:
                    self.add_hit(SearchHit(
                        "ERROR", item, "", "", f"{type(exc).__name__}: {exc}", (), "ERROR"
                    ), search_id)
                finally:
                    elapsed = time.perf_counter() - started
                    ext = os.path.splitext(item)[1].lower() or "other"
                    with self.stats_lock:
                        self.stats.completed_files += 1
                        self.stats.file_type_seconds[ext] += elapsed
                        self.stats.active_files.pop(thread_name, None)
            finally:
                work_q.task_done()

    def finish_search(self, search_id: int, stopped: bool):
        if search_id != self._search_id:
            return
        with self.stats_lock:
            self.stats.finished_at = time.perf_counter()
            self.stats.phase = "中断" if stopped else "完了"
        self._search_running = False
        self._close_autosave()

    # -------------------------------------------------------------- processors
    def process_file(self, path: str, search_id: int, options: SearchOptions, matcher: SearchMatcher):
        self.pause_point()
        if options.exclude_query and matcher.contains_excluded(self.path_candidates(path, options.roots)):
            return

        if options.search_path_names:
            candidates = self.path_candidates(path, options.roots)
            matched = matcher.evaluate(candidates)
            if matched:
                self.add_hit(SearchHit(
                    "Path", path, "ファイル名/フォルダ名", "",
                    self.relative_path(path, options.roots), matched,
                ), search_id)

        ext = os.path.splitext(path)[1].lower()

        # Unknown proprietary extensions are probed here, not during directory enumeration.
        # This avoids needlessly hydrating OneDrive/SPO Files-On-Demand placeholders before
        # real search work has even begun.
        if ext not in ALL_SUPPORTED_EXTS:
            if not (options.include_text and options.include_unknown_text and looks_like_text_file(path)):
                return
            self.search_text(path, path, search_id, matcher, file_type="Text(推定)")
            return

        # For PDFs, serialize staging as well as parsing. This prevents four
        # workers from simultaneously creating large temporary PDF copies.
        if ext in PDF_EXTS:
            with self._pdf_semaphore:
                with local_parse_path(path, options) as parse_path:
                    self.search_pdf(parse_path, path, search_id, matcher)
            return

        with local_parse_path(path, options) as parse_path:
            if ext in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
                self.search_xlsx(parse_path, path, search_id, options, matcher)
            elif ext == ".xls":
                self.search_xls(parse_path, path, search_id, matcher)
            elif ext == ".xlsb":
                self.search_xlsb(parse_path, path, search_id, matcher)
            elif ext in WORD_EXTS:
                self.search_word(parse_path, path, search_id, matcher)
            elif ext in POWERPOINT_EXTS:
                self.search_powerpoint(parse_path, path, search_id, matcher)
            elif ext in TEXT_EXTS:
                self.search_text(parse_path, path, search_id, matcher)

    def search_xlsx(self, parse_path: str, original_path: str, search_id: int, options: SearchOptions, matcher: SearchMatcher):
        if load_workbook is None:
            self.add_hit(SearchHit("ERROR", original_path, "", "", "openpyxl がありません: pip install openpyxl", (), "ERROR"), search_id)
            return

        evidence: list[TextEvidence] = []
        data_wb = None
        formula_wb = None
        try:
            data_wb = load_workbook(parse_path, read_only=True, data_only=True, keep_links=False)
            if options.search_formula:
                try:
                    formula_wb = load_workbook(parse_path, read_only=True, data_only=False, keep_links=False)
                except Exception:
                    formula_wb = None

            for data_ws in data_wb.worksheets:
                self.pause_point()
                formula_ws = formula_wb[data_ws.title] if formula_wb is not None and data_ws.title in formula_wb.sheetnames else None
                data_rows = data_ws.iter_rows()
                formula_rows = formula_ws.iter_rows() if formula_ws is not None else iter(())

                for row_no, (data_row, formula_row) in enumerate(itertools.zip_longest(data_rows, formula_rows, fillvalue=()), 1):
                    if row_no % 50 == 0:
                        self.pause_point()
                    max_cols = max(len(data_row), len(formula_row))
                    for col_index in range(max_cols):
                        dcell = data_row[col_index] if col_index < len(data_row) else None
                        fcell = formula_row[col_index] if col_index < len(formula_row) else None
                        value = getattr(dcell, "value", None)
                        formula = getattr(fcell, "value", None)
                        if value is None and formula is None:
                            continue
                        number_format = getattr(dcell, "number_format", "") if dcell is not None else ""
                        candidates = excel_cell_candidates(value, number_format)
                        if isinstance(formula, str) and formula.startswith("="):
                            candidates.append(formula)
                        if not candidates:
                            continue
                        coord = getattr(dcell, "coordinate", None) or getattr(fcell, "coordinate", None)
                        if not coord:
                            coord = f"{get_column_letter(col_index + 1)}{row_no}" if get_column_letter else f"R{row_no}C{col_index + 1}"
                        shown = candidates[0]
                        if isinstance(formula, str) and formula.startswith("=") and formula not in str(shown):
                            shown = f"{shown}   [式: {formula}]" if shown else formula
                        terms = matcher.unit_terms(candidates)
                        if terms:
                            evidence.append(TextEvidence(
                                "Excel", original_path, data_ws.title, coord, matcher.snippet(shown), terms
                            ))
            self.emit_file_evidence(evidence, search_id, matcher)
        finally:
            for wb in (data_wb, formula_wb):
                if wb is not None:
                    try:
                        wb.close()
                    except Exception:
                        pass


    def search_xls(self, parse_path: str, original_path: str, search_id: int, matcher: SearchMatcher):
        if xlrd is None:
            self.add_hit(SearchHit("ERROR", original_path, "", "", "xlrd がありません: pip install xlrd", (), "ERROR"), search_id)
            return
        evidence: list[TextEvidence] = []
        book = xlrd.open_workbook(parse_path, on_demand=True)
        try:
            for sheet in book.sheets():
                for r in range(sheet.nrows):
                    if r % 50 == 0:
                        self.pause_point()
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        value = cell.value
                        if value in (None, ""):
                            continue
                        try:
                            if cell.ctype == getattr(xlrd, "XL_CELL_DATE", -999):
                                value = xlrd.xldate_as_datetime(value, book.datemode)
                        except Exception:
                            pass
                        candidates = excel_cell_candidates(value)
                        terms = matcher.unit_terms(candidates)
                        if terms:
                            coord = f"{get_column_letter(c + 1)}{r + 1}" if get_column_letter else f"R{r + 1}C{c + 1}"
                            evidence.append(TextEvidence(
                                "Excel", original_path, sheet.name, coord, matcher.snippet(candidates[0]), terms
                            ))
            self.emit_file_evidence(evidence, search_id, matcher)
        finally:
            book.release_resources()


    def search_xlsb(self, parse_path: str, original_path: str, search_id: int, matcher: SearchMatcher):
        if open_xlsb_workbook is None:
            self.add_hit(SearchHit("ERROR", original_path, "", "", "pyxlsb がありません: pip install pyxlsb", (), "ERROR"), search_id)
            return
        evidence: list[TextEvidence] = []
        with open_xlsb_workbook(parse_path) as wb:
            for sheet_name in wb.sheets:
                with wb.get_sheet(sheet_name) as sheet:
                    for r, row in enumerate(sheet.rows(), 1):
                        if r % 50 == 0:
                            self.pause_point()
                        for c, cell in enumerate(row, 1):
                            value = cell.v
                            if value is None:
                                continue
                            candidates = excel_cell_candidates(value)
                            terms = matcher.unit_terms(candidates)
                            if terms:
                                coord = f"{get_column_letter(c)}{r}" if get_column_letter else f"R{r}C{c}"
                                evidence.append(TextEvidence(
                                    "Excel", original_path, sheet_name, coord, matcher.snippet(candidates[0]), terms
                                ))
        self.emit_file_evidence(evidence, search_id, matcher)


    def search_pdf(self, parse_path: str, original_path: str, search_id: int, matcher: SearchMatcher):
        if PdfReader is None:
            self.add_hit(SearchHit("ERROR", original_path, "", "", "pypdf がありません: pip install pypdf", (), "ERROR"), search_id)
            return
        evidence: list[TextEvidence] = []
        reader = PdfReader(parse_path, strict=False)
        try:
            if getattr(reader, "is_encrypted", False):
                try:
                    if not reader.decrypt(""):
                        self.add_hit(SearchHit("PDF", original_path, "", "", "暗号化PDFのため検索できません", (), "INFO"), search_id)
                        return
                except Exception:
                    self.add_hit(SearchHit("PDF", original_path, "", "", "暗号化PDFのため検索できません", (), "INFO"), search_id)
                    return
        except Exception:
            pass

        ocr_mode = (self._active_options.pdf_ocr_mode if self._active_options else "OFF") or "OFF"
        wants_ocr = ocr_mode != "OFF"
        ocr_ready = bool(self._ocr_ready and pdfium is not None and pytesseract is not None)
        ocr_doc = None
        if wants_ocr and not ocr_ready:
            self.add_hit(SearchHit(
                "PDF", original_path, "", "",
                f"OCR未使用: {self._ocr_status}。通常PDF文字だけ検索しました。",
                (), "INFO"
            ), search_id)
        elif wants_ocr:
            try:
                ocr_doc = pdfium.PdfDocument(parse_path)
            except Exception as exc:
                self.add_hit(SearchHit("ERROR", original_path, "", "", f"OCR用PDF初期化エラー: {exc}", (), "ERROR"), search_id)
                ocr_ready = False

        try:
            for i, page in enumerate(reader.pages, 1):
                self.pause_point()
                native_text = ""
                try:
                    native_text = page.extract_text() or ""
                except Exception as exc:
                    self.add_hit(SearchHit("ERROR", original_path, f"{i}ページ", "", f"PDF抽出エラー: {exc}", (), "ERROR"), search_id)

                combined_text = native_text
                ocr_text = ""
                should_ocr = False
                if ocr_ready and ocr_mode == "全ページ":
                    should_ocr = True
                elif ocr_ready and ocr_mode == "自動":
                    # OCR only pages whose native text layer is absent or very small.
                    # This catches ordinary scanned PDFs without turning every digital PDF
                    # into an expensive image-recognition job.
                    should_ocr = len(native_text.strip()) < 40

                if should_ocr and ocr_doc is not None:
                    try:
                        ocr_text = self.ocr_pdfium_page(ocr_doc, i - 1)
                        if ocr_text:
                            combined_text = (native_text + "\n" + ocr_text).strip()
                    except RuntimeError as exc:
                        self.add_hit(SearchHit("INFO", original_path, f"{i}ページ", "", f"OCRタイムアウト/中断: {exc}", (), "INFO"), search_id)
                    except Exception as exc:
                        self.add_hit(SearchHit("ERROR", original_path, f"{i}ページ", "", f"OCRエラー: {type(exc).__name__}: {exc}", (), "ERROR"), search_id)

                if not combined_text:
                    continue
                terms = matcher.unit_terms([combined_text])
                if terms:
                    source = "OCR" if ocr_text and not native_text.strip() else ("PDF文字+OCR" if ocr_text else "PDF文字")
                    evidence.append(TextEvidence(
                        "PDF", original_path, f"{i}ページ [{source}]", "", matcher.snippet(combined_text), terms
                    ))
            self.emit_file_evidence(evidence, search_id, matcher)
        finally:
            if ocr_doc is not None:
                try:
                    ocr_doc.close()
                except Exception:
                    pass

    def ocr_pdfium_page(self, pdf, page_index: int) -> str:
        if not self._ocr_ready or pytesseract is None:
            return ""
        page = pdf[page_index]
        bitmap = None
        image = None
        try:
            # 2x is a reasonable speed/accuracy compromise for document search.
            bitmap = page.render(scale=2.0)
            image = bitmap.to_pil()
            return pytesseract.image_to_string(
                image, lang=self._ocr_lang or "eng", timeout=OCR_PAGE_TIMEOUT_SECONDS
            ) or ""
        finally:
            try:
                if image is not None:
                    image.close()
            except Exception:
                pass
            try:
                if bitmap is not None and hasattr(bitmap, "close"):
                    bitmap.close()
            except Exception:
                pass
            try:
                if hasattr(page, "close"):
                    page.close()
            except Exception:
                pass


    def search_word(self, parse_path: str, original_path: str, search_id: int, matcher: SearchMatcher):
        if Document is None:
            self.add_hit(SearchHit("ERROR", original_path, "", "", "python-docx がありません: pip install python-docx", (), "ERROR"), search_id)
            return
        evidence: list[TextEvidence] = []
        doc = Document(parse_path)
        for i, para in enumerate(doc.paragraphs, 1):
            if i % 50 == 0:
                self.pause_point()
            text = para.text or ""
            if text:
                terms = matcher.unit_terms([text])
                if terms:
                    evidence.append(TextEvidence("Word", original_path, f"段落 {i}", "", matcher.snippet(text), terms))

        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                self.pause_point()
                for ci, cell in enumerate(row.cells, 1):
                    text = cell.text or ""
                    if text:
                        terms = matcher.unit_terms([text])
                        if terms:
                            evidence.append(TextEvidence(
                                "Word", original_path, f"表{ti} 行{ri} 列{ci}", "", matcher.snippet(text), terms
                            ))
        self.emit_file_evidence(evidence, search_id, matcher)


    def search_powerpoint(self, parse_path: str, original_path: str, search_id: int, matcher: SearchMatcher):
        if Presentation is None:
            self.add_hit(SearchHit("ERROR", original_path, "", "", "python-pptx がありません: pip install python-pptx", (), "ERROR"), search_id)
            return
        evidence: list[TextEvidence] = []
        prs = Presentation(parse_path)
        for si, slide in enumerate(prs.slides, 1):
            self.pause_point()
            for shape_index, shape in enumerate(slide.shapes, 1):
                texts: list[tuple[str, str]] = []
                try:
                    if getattr(shape, "has_text_frame", False):
                        text = shape.text or ""
                        if text:
                            texts.append((f"スライド{si} 図形{shape_index}", text))
                except Exception:
                    pass
                if getattr(shape, "has_table", False):
                    try:
                        for ri, row in enumerate(shape.table.rows, 1):
                            for ci, cell in enumerate(row.cells, 1):
                                text = cell.text or ""
                                if text:
                                    texts.append((f"スライド{si} 表 行{ri} 列{ci}", text))
                    except Exception:
                        pass
                for place, text in texts:
                    terms = matcher.unit_terms([text])
                    if terms:
                        evidence.append(TextEvidence(
                            "PowerPoint", original_path, place, "", matcher.snippet(text), terms
                        ))

            try:
                if slide.has_notes_slide:
                    text = slide.notes_slide.notes_text_frame.text or ""
                    terms = matcher.unit_terms([text]) if text else ()
                    if terms:
                        evidence.append(TextEvidence(
                            "PowerPoint", original_path, f"スライド{si} ノート", "", matcher.snippet(text), terms
                        ))
            except Exception:
                pass
        self.emit_file_evidence(evidence, search_id, matcher)


    def search_text(self, parse_path: str, original_path: str, search_id: int, matcher: SearchMatcher, file_type: str = "Text"):
        evidence: list[TextEvidence] = []
        try:
            with open(parse_path, "rb") as raw:
                encoding = detect_text_encoding(raw)
                wrapper = io.TextIOWrapper(raw, encoding=encoding, errors="replace", newline=None)
                for line_no, line in enumerate(wrapper, 1):
                    if line_no % 200 == 0:
                        self.pause_point()
                    if not line.strip():
                        continue
                    terms = matcher.unit_terms([line])
                    if terms:
                        evidence.append(TextEvidence(
                            file_type, original_path, f"{line_no}行目", "", matcher.snippet(line), terms
                        ))
            self.emit_file_evidence(evidence, search_id, matcher)
        except Exception as exc:
            self.add_hit(SearchHit("ERROR", original_path, "", "", f"テキスト読込エラー: {exc}", (), "ERROR"), search_id)


    def emit_file_evidence(self, evidence: list[TextEvidence], search_id: int, matcher: SearchMatcher):
        """Evaluate one whole file, then emit the proof rows that made it match.

        This is the main final-version behavior: A&B can match even when A and B
        live in different cells, sheets, pages, slides, or lines of the same file.
        """
        if not evidence:
            return
        present_terms: set[str] = set()
        for item in evidence:
            present_terms.update(item.matched_terms)
        matched_groups = matcher.file_displays(present_terms)
        if not matched_groups:
            return
        wanted_terms = matcher.terms_for_displays(matched_groups)

        # Summary row is useful when AND terms are spread across distant places.
        first = evidence[0]
        if matcher.threshold_min is not None or any("&" in group for group in matched_groups) or len(matched_groups) > 1:
            self.add_hit(SearchHit(
                first.file_type, first.path, "ファイル全体一致", "",
                " / ".join(matched_groups), matched_groups
            ), search_id)

        emitted = 0
        for item in evidence:
            terms = tuple(t for t in item.matched_terms if t in wanted_terms)
            if not terms:
                continue
            self.add_hit(SearchHit(
                item.file_type, item.path, item.place, item.cell, item.value, terms
            ), search_id)
            emitted += 1
            if emitted >= 5000:
                # Keep pathological "A" searches from creating huge proof lists per file.
                break

    def add_evidence_if_match(
        self,
        evidence: list[TextEvidence],
        matcher: SearchMatcher,
        file_type: str,
        original_path: str,
        place: str,
        cell: str,
        raw_text: object,
        shown_text: object | None = None,
    ):
        terms = matcher.unit_terms([raw_text])
        if not terms:
            return
        shown = raw_text if shown_text is None else shown_text
        evidence.append(TextEvidence(
            file_type, original_path, place, cell, matcher.snippet(shown), terms
        ))

    def selected_tree(self):
        focus = self.root.focus_get()
        if focus is self.keyword_tree:
            return self.keyword_tree, self._keyword_tree_map
        return self.result_tree, self._result_tree_map

    def normalized_confirmed_path(self, path: str) -> str:
        try:
            return os.path.normcase(os.path.normpath(str(path)))
        except Exception:
            return str(path)

    def is_confirmed_path(self, path: str) -> bool:
        return self.normalized_confirmed_path(path) in self.confirmed_files

    def confirmed_label(self, path: str) -> str:
        return "済" if self.is_confirmed_path(path) else ""


    # ------------------------------------------------------------- hit storage
    def add_hit(self, hit: SearchHit, search_id: int):
        if search_id != self._search_id:
            return

        new_hit_file = False
        with self.results_lock:
            # Keep interactive RAM use bounded. Full results continue to the
            # autosave CSV even after this limit is reached.
            store_in_memory = len(self.results) < MAX_IN_MEMORY_HITS
            index = len(self.results) if store_in_memory else -1
            if store_in_memory:
                self.results.append(hit)
            else:
                self._memory_hit_limit_reached = True

            if hit.severity == "HIT":
                new_hit_file = hit.path not in self.hit_files
                self.file_hit_counts[hit.path] += 1
                self.type_hit_counts[hit.file_type] += 1
                self.hit_files.add(hit.path)
                for keyword in hit.matched_keywords:
                    self.keyword_hit_counts[keyword] += 1
                    if store_in_memory:
                        self.keyword_results[keyword].append(index)

        with self.stats_lock:
            if hit.severity == "HIT":
                self.stats.hit_count += 1
                if new_hit_file:
                    self.stats.hit_files += 1
            elif hit.severity == "ERROR":
                self.stats.errors += 1

        self._write_autosave(hit)

    # --------------------------------------------------------------- UI sync
    def _schedule_ui_refresh(self):
        if self._ui_job is None:
            self._ui_job = self.root.after(UI_REFRESH_MS, self._refresh_ui)

    def _refresh_ui(self):
        self._ui_job = None
        try:
            with self.results_lock:
                total_results = len(self.results)
                display_end = min(total_results, self._display_limit)
                new_hits = list(enumerate(self.results[self._rendered_result_count:display_end], self._rendered_result_count))

            for index, hit in new_hits:
                self.insert_result_row(index, hit)
            self._rendered_result_count = display_end
            self.display_info_var.set(f"表示 {self._rendered_result_count} / {total_results}")

            now = time.monotonic()
            if now - self._last_keyword_refresh >= 1.0:
                self.refresh_keyword_list()
                self._last_keyword_refresh = now

            with self.stats_lock:
                stats = RuntimeStats(
                    total_files=self.stats.total_files,
                    completed_files=self.stats.completed_files,
                    hit_count=self.stats.hit_count,
                    hit_files=self.stats.hit_files,
                    errors=self.stats.errors,
                    started_at=self.stats.started_at,
                    finished_at=self.stats.finished_at,
                    phase=self.stats.phase,
                    last_file=self.stats.last_file,
                    active_files=dict(self.stats.active_files),
                    file_type_seconds=dict(self.stats.file_type_seconds),
                )

            self.progress_bar.config(maximum=max(stats.total_files, 1))
            self.progress_var.set(min(stats.completed_files, max(stats.total_files, 1)))
            elapsed_end = stats.finished_at or time.perf_counter()
            elapsed = max(0.0, elapsed_end - stats.started_at) if stats.started_at else 0.0
            if stats.phase == "列挙＋検索中":
                progress_text = f"処理 {stats.completed_files} / 発見 {stats.total_files} ファイル"
            else:
                progress_text = f"{stats.completed_files} / {stats.total_files} ファイル"
            self.summary_var.set(
                f"{progress_text}   {stats.hit_count} ヒット / {stats.hit_files}資料   {elapsed:.1f}s"
            )

            if self._search_running:
                if not self.pause_event.is_set():
                    self.status_var.set("一時停止中...")
                else:
                    active = [os.path.basename(v[0]) for v in stats.active_files.values()]
                    suffix = " / ".join(active[:3])
                    self.status_var.set(f"{stats.phase}: {suffix}" if suffix else stats.phase)
            else:
                if stats.phase in {"完了", "中断"}:
                    info = f"{stats.phase}: {stats.hit_count}ヒット / {stats.hit_files}資料 / エラー{stats.errors}件 / {elapsed:.1f}秒"
                    if self._autosave_path:
                        info += f" / 自動退避: {self._autosave_path}"
                    if self._memory_hit_limit_reached:
                        info += f" / 画面保持は先頭{MAX_IN_MEMORY_HITS:,}件（全件は自動退避CSV）"
                    if self._low_disk_stop:
                        info += " / TEMP残容量保護のため中断"
                    self.status_var.set(info)
                    self.search_button.config(state=tk.NORMAL)
                    self.pause_button.config(state=tk.DISABLED, text="一時停止")
                    self.stop_button.config(state=tk.DISABLED)
        finally:
            self._schedule_ui_refresh()

    def insert_result_row(self, index: int, hit: SearchHit):
        tags = []
        if hit.severity == "ERROR":
            tags.append("error")
        elif hit.severity == "INFO":
            tags.append("info")
        if self.is_confirmed_path(hit.path):
            tags.append("confirmed")
        iid = self.result_tree.insert(
            "", tk.END,
            values=(
                self.confirmed_label(hit.path), os.path.basename(hit.path), ", ".join(hit.matched_keywords),
                hit.place, hit.cell, hit.value, hit.file_type, hit.path,
            ),
            tags=tuple(tags),
        )
        self.result_tree.tag_configure("error", foreground="#a00000")
        self.result_tree.tag_configure("info", foreground="#355c7d")
        self.result_tree.tag_configure("confirmed", background="#e7e7e7")
        self._result_tree_map[iid] = index

    def clear_ui_results(self):
        """Clear visible result widgets without deleting the underlying result list.

        This must also reset _rendered_result_count. Otherwise a rebuild after
        sorting or a confirmed-state operation can leave the Results tab empty
        because the refresh loop thinks every row has already been drawn.
        """
        for tree in (getattr(self, "result_tree", None), getattr(self, "keyword_tree", None)):
            if tree is not None:
                tree.delete(*tree.get_children())
        if hasattr(self, "keyword_listbox"):
            self.keyword_listbox.delete(0, tk.END)
        self._result_tree_map.clear()
        self._keyword_tree_map.clear()
        self._keyword_list_keys.clear()
        self._rendered_result_count = 0
        total = len(self.results) if hasattr(self, "results") else 0
        self.display_info_var.set(f"表示 0 / {total}")


    def result_sort_value(self, hit: SearchHit, column: str):
        if column == "confirmed":
            return 0 if self.is_confirmed_path(hit.path) else 1
        if column == "file":
            return os.path.basename(hit.path).casefold()
        if column == "keywords":
            return ", ".join(hit.matched_keywords).casefold()
        if column == "place":
            return hit.place.casefold()
        if column == "cell":
            return hit.cell.casefold()
        if column == "value":
            return hit.value.casefold()
        if column == "type":
            return hit.file_type.casefold()
        if column == "path":
            return hit.path.casefold()
        return ""

    def sort_results_by(self, column: str):
        with self.results_lock:
            if self._sort_column == column:
                self._sort_desc = not self._sort_desc
            else:
                self._sort_column = column
                self._sort_desc = False
            self.results.sort(key=lambda h: self.result_sort_value(h, column), reverse=self._sort_desc)
            self.rebuild_keyword_indexes_locked()
        self.clear_ui_results()
        self._display_limit = max(self._display_limit, INITIAL_DISPLAY_ROWS)
        self.refresh_keyword_list(force=True)
        self._schedule_ui_refresh()

    def sort_keyword_results_by(self, column: str):
        if self._keyword_sort_column == column:
            self._keyword_sort_desc = not self._keyword_sort_desc
        else:
            self._keyword_sort_column = column
            self._keyword_sort_desc = False
        self.on_keyword_select()

    def rebuild_keyword_indexes_locked(self):
        self.keyword_results = defaultdict(list)
        self.keyword_hit_counts = defaultdict(int)
        self.file_hit_counts = defaultdict(int)
        self.type_hit_counts = defaultdict(int)
        self.hit_files = set()
        for index, hit in enumerate(self.results):
            if hit.severity != "HIT":
                continue
            self.file_hit_counts[hit.path] += 1
            self.type_hit_counts[hit.file_type] += 1
            self.hit_files.add(hit.path)
            for keyword in hit.matched_keywords:
                self.keyword_hit_counts[keyword] += 1
                self.keyword_results[keyword].append(index)


    def load_more_results(self):
        self._display_limit += LOAD_MORE_ROWS
        self._schedule_ui_refresh()

    def refresh_keyword_list(self, force: bool = False):
        with self.results_lock:
            items = sorted(self.keyword_hit_counts.items(), key=lambda x: (-x[1], x[0]))
        selected_key = None
        selection = self.keyword_listbox.curselection()
        if selection and selection[0] < len(self._keyword_list_keys):
            selected_key = self._keyword_list_keys[selection[0]]

        self.keyword_listbox.delete(0, tk.END)
        self._keyword_list_keys = []
        restore_index = None
        for i, (key, count) in enumerate(items):
            self._keyword_list_keys.append(key)
            self.keyword_listbox.insert(tk.END, f"{key} ({count}件)")
            if key == selected_key:
                restore_index = i
        if restore_index is not None:
            self.keyword_listbox.selection_set(restore_index)

    def on_keyword_select(self, _event=None):
        selection = self.keyword_listbox.curselection()
        if not selection or selection[0] >= len(self._keyword_list_keys):
            return
        keyword = self._keyword_list_keys[selection[0]]
        self.keyword_tree.delete(*self.keyword_tree.get_children())
        self._keyword_tree_map.clear()
        with self.results_lock:
            indices = list(self.keyword_results.get(keyword, []))[:KEYWORD_DETAIL_ROWS]
            rows = [(i, self.results[i]) for i in indices if i < len(self.results)]
        if self._keyword_sort_column:
            rows.sort(
                key=lambda pair: self.result_sort_value(pair[1], self._keyword_sort_column),
                reverse=self._keyword_sort_desc,
            )
        for index, hit in rows:
            tags = []
            if hit.severity == "ERROR":
                tags.append("error")
            elif hit.severity == "INFO":
                tags.append("info")
            if self.is_confirmed_path(hit.path):
                tags.append("confirmed")
            iid = self.keyword_tree.insert("", tk.END, values=(
                self.confirmed_label(hit.path), os.path.basename(hit.path), hit.place, hit.cell,
                hit.value, hit.file_type, hit.path,
            ), tags=tuple(tags))
            self.keyword_tree.tag_configure("error", foreground="#a00000")
            self.keyword_tree.tag_configure("info", foreground="#355c7d")
            self.keyword_tree.tag_configure("confirmed", background="#e7e7e7")
            self._keyword_tree_map[iid] = index

    # ------------------------------------------------------------- path utils
    @staticmethod
    def relative_path(path: str, roots: tuple[str, ...]) -> str:
        abs_path = os.path.abspath(path)
        candidates = sorted((os.path.abspath(r) for r in roots), key=len, reverse=True)
        for root in candidates:
            try:
                rel = os.path.relpath(abs_path, root)
                if rel != ".." and not rel.startswith(".." + os.sep):
                    return rel
            except Exception:
                pass
        return os.path.basename(path)

    def path_candidates(self, path: str, roots: tuple[str, ...]) -> list[str]:
        rel = self.relative_path(path, roots)
        name = os.path.basename(path)
        directory = os.path.dirname(rel)
        values = [rel, name]
        if directory:
            values.append(directory)
            values.extend(part for part in re.split(r"[\\/]+", directory) if part)
        return list(dict.fromkeys(v for v in values if v))

    # -------------------------------------------------------------- autosave
    def _start_autosave(self, query: str):
        self._close_autosave()
        try:
            stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
            self._autosave_path = os.path.join(tempfile.gettempdir(), f"file_search_v2_{stamp}.csv")
            self._autosave_fp = open(self._autosave_path, "w", newline="", encoding="utf-8-sig")
            self._autosave_writer = csv.writer(self._autosave_fp)
            self._autosave_writer.writerow([
                "種類", "ファイルパス", "場所", "セル", "内容", "マッチキーワード", "状態", "検索文字",
            ])
            self._autosave_fp.flush()
            self._autosave_count = 0
        except Exception:
            self._autosave_path = None
            self._autosave_fp = None
            self._autosave_writer = None

    def _write_autosave(self, hit: SearchHit):
        if self._autosave_writer is None:
            return
        try:
            with self._autosave_lock:
                self._autosave_writer.writerow([
                    hit.file_type, hit.path, hit.place, hit.cell, hit.value,
                    ", ".join(hit.matched_keywords), hit.severity,
                    self._active_options.query if self._active_options else "",
                ])
                self._autosave_count += 1
                if self._autosave_count % AUTOSAVE_FLUSH_EVERY == 0:
                    self._autosave_fp.flush()

                # Periodically protect the system drive from an unexpectedly
                # huge hit set. Stop cleanly before TEMP consumes the last few GB.
                if self._autosave_count % 10_000 == 0:
                    try:
                        free_bytes = shutil.disk_usage(tempfile.gettempdir()).free
                        if free_bytes < AUTOSAVE_LOW_DISK_STOP_BYTES:
                            self._low_disk_stop = True
                            self.stop_event.set()
                            self.pause_event.set()
                    except Exception:
                        pass
        except Exception:
            pass

    def _close_autosave(self):
        with self._autosave_lock:
            try:
                if self._autosave_fp is not None:
                    self._autosave_fp.flush()
                    self._autosave_fp.close()
            except Exception:
                pass
            finally:
                self._autosave_fp = None
                self._autosave_writer = None

    def save_csv(self):
        with self.results_lock:
            rows = list(self.results)
        if not rows:
            messagebox.showinfo("確認", "保存する検索結果がありません。")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv", filetypes=[("CSV", "*.csv"), ("すべて", "*.*")],
            initialfile="search_results.csv",
        )
        if not path:
            return
        try:
            if self._memory_hit_limit_reached and self._autosave_path and os.path.isfile(self._autosave_path):
                if self._search_running:
                    messagebox.showwarning("確認", "大量ヒット検索中です。全件CSVは検索完了後に保存してください。")
                    return
                shutil.copyfile(self._autosave_path, path)
            else:
                with open(path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    writer.writerow(["種類", "ファイルパス", "場所", "セル", "内容", "マッチキーワード", "状態"])
                    for hit in rows:
                        writer.writerow([
                            hit.file_type, hit.path, hit.place, hit.cell, hit.value,
                            ", ".join(hit.matched_keywords), hit.severity,
                        ])
            messagebox.showinfo("完了", f"保存しました。\n{path}")
        except Exception as exc:
            messagebox.showerror("エラー", f"CSV保存に失敗しました。\n{exc}")

    # ------------------------------------------------------------ context menu
    def selected_hit(self) -> SearchHit | None:
        widget = self.root.focus_get()
        if widget not in (self.result_tree, self.keyword_tree):
            return None
        selection = widget.selection()
        if not selection:
            return None
        iid = selection[0]
        index = self._result_tree_map.get(iid) if widget is self.result_tree else self._keyword_tree_map.get(iid)
        if index is None:
            return None
        with self.results_lock:
            if 0 <= index < len(self.results):
                return self.results[index]
        return None

    def show_context_menu(self, event):
        tree = event.widget
        iid = tree.identify_row(event.y)
        if iid:
            tree.selection_set(iid)
            tree.focus(iid)
            tree.focus_set()
            self.context_menu.tk_popup(event.x_root, event.y_root)

    def open_selected_smart(self):
        hit = self.selected_hit()
        if hit and hit.file_type == "Excel" and hit.cell:
            self.open_selected_excel_cell()
        else:
            self.open_selected_file()

    def open_selected_file(self):
        hit = self.selected_hit()
        if not hit or not hit.path:
            return
        try:
            os.startfile(os.path.normpath(hit.path))
        except Exception as exc:
            messagebox.showerror("エラー", f"ファイルを開けません。\n{exc}")

    def open_selected_folder(self):
        hit = self.selected_hit()
        if not hit or not hit.path:
            return
        path = os.path.normpath(hit.path)
        try:
            if os.name == "nt":
                subprocess.Popen(["explorer", "/select,", path])
            else:
                subprocess.Popen(["xdg-open", os.path.dirname(path)])
        except Exception as exc:
            messagebox.showerror("エラー", f"フォルダを開けません。\n{exc}")

    def open_selected_excel_cell(self):
        hit = self.selected_hit()
        if not hit or hit.file_type != "Excel" or not hit.path:
            messagebox.showinfo("確認", "Excelの検索結果を選択してください。")
            return
        if win32 is None:
            self.open_selected_file()
            return
        try:
            excel = win32.Dispatch("Excel.Application")
            excel.Visible = True
            wb = excel.Workbooks.Open(os.path.abspath(hit.path))
            if hit.place and hit.place in [ws.Name for ws in wb.Worksheets]:
                ws = wb.Worksheets(hit.place)
                ws.Activate()
                if hit.cell:
                    ws.Range(hit.cell).Select()
        except Exception as exc:
            messagebox.showwarning("確認", f"該当セルを直接開けませんでした。\n{exc}\n\nファイルを開きます。")
            self.open_selected_file()

    def copy_selected_value(self):
        hit = self.selected_hit()
        if hit:
            self.copy_to_clipboard(hit.value)

    def copy_selected_path(self):
        hit = self.selected_hit()
        if hit:
            self.copy_to_clipboard(hit.path)

    def copy_to_clipboard(self, text: str):
        self.root.clipboard_clear()
        self.root.clipboard_append(str(text))
        self.root.update_idletasks()

    # ---------------------------------------------------------------- closing

    def toggle_confirmed_state(self):
        hit = self.selected_hit()
        if not hit or not hit.path:
            return
        norm = self.normalized_confirmed_path(hit.path)
        if norm in self.confirmed_files:
            self.confirmed_files.remove(norm)
        else:
            self.confirmed_files.add(norm)
        self.save_settings()
        self.refresh_confirmed_state_for_path(hit.path)
        self._schedule_ui_refresh()
        self.status_var.set("確認済み状態を更新しました")

    def refresh_confirmed_state_for_path(self, path: str):
        norm = self.normalized_confirmed_path(path)
        for tree, mapping in ((self.result_tree, self._result_tree_map), (self.keyword_tree, self._keyword_tree_map)):
            for iid, index in list(mapping.items()):
                try:
                    with self.results_lock:
                        hit = self.results[index]
                    if self.normalized_confirmed_path(hit.path) != norm:
                        continue
                    values = list(tree.item(iid, "values"))
                    if values:
                        values[0] = self.confirmed_label(hit.path)
                        tree.item(iid, values=values)
                    tags = list(tree.item(iid, "tags"))
                    if self.is_confirmed_path(hit.path):
                        if "confirmed" not in tags:
                            tags.append("confirmed")
                    else:
                        tags = [t for t in tags if t != "confirmed"]
                    tree.item(iid, tags=tuple(tags))
                except Exception:
                    pass


    def on_close(self):
        self.stop_event.set()
        self.pause_event.set()
        self.add_history(self.keyword_var.get())
        self.save_settings()
        self._close_autosave()
        try:
            if self._ui_job is not None:
                self.root.after_cancel(self._ui_job)
        except Exception:
            pass
        self.root.destroy()


if __name__ == "__main__":
    root = tk.Tk()
    app = SearchApp(root)
    root.mainloop()
