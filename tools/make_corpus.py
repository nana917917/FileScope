"""Build a small fixture corpus with the document formats FileScope supports.

Used by the tests and by ``scripts/benchmark.py`` so both exercise the same
real files instead of mocks. Everything is generated locally; no sample
documents are committed.
"""

from __future__ import annotations

import os
import zipfile
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Corpus:
    root: Path
    files: dict[str, Path] = field(default_factory=dict)

    def path(self, name: str) -> Path:
        return self.files[name]


# --------------------------------------------------------------------- PDF


def write_text_pdf(path: Path, pages: list[str]) -> None:
    """Write a minimal PDF whose pages carry a real text layer."""
    objects: list[bytes] = []
    page_ids = [3 + index * 2 for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + len(pages) * 2
    for index, text in enumerate(pages):
        content_id = page_ids[index] + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
            ).encode()
        )
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        stream = f"BT /F1 18 Tf 72 760 Td ({escaped}) Tj ET".encode("latin-1", "replace")
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    path.write_bytes(bytes(out))


def write_image_pdf(path: Path, text: str) -> None:
    """Write an image-only PDF (no text layer) for OCR tests."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (1000, 320), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 130), text, fill="black")
    image.save(path, "PDF", resolution=150.0)


# ------------------------------------------------------------------ Office


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "評価"
    sheet["A1"] = "AAA"
    sheet["B2"] = "=SUM(1,2)"
    sheet["A3"] = "耐久"
    sheet["A3"].comment = Comment("別シート確認のこと", "tester")
    for index in range(2, 9):
        workbook.create_sheet(f"Sheet{index}")
    workbook["Sheet8"]["D52"] = "BBB"
    workbook.defined_names.add(__import__("openpyxl").workbook.defined_name.DefinedName("評価範囲", attr_text="評価!$A$1"))
    workbook.save(path)


def write_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("AAA 落下試験の結果")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "BBB"
    table.cell(0, 1).text = "評価"
    document.sections[0].header.paragraphs[0].text = "社外秘ヘッダ"
    document.save(path)


def write_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "AAA 耐久レビュー"
    slide.notes_slide.notes_text_frame.text = "BBB ノート"
    hidden = presentation.slides.add_slide(presentation.slide_layouts[5])
    hidden.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "CCC 非表示"
    hidden._element.set("show", "0")
    presentation.save(path)


def write_zip(path: Path, entries: dict[str, str]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, text in entries.items():
            archive.writestr(name, text)


# ------------------------------------------------------------------ corpus


def build_corpus(root: str | os.PathLike[str], *, with_ocr_image: bool = True) -> Corpus:
    root_path = Path(root)
    root_path.mkdir(parents=True, exist_ok=True)
    corpus = Corpus(root=root_path)

    def add(name: str, builder) -> None:
        target = root_path / name
        builder(target)
        corpus.files[name] = target

    add("text_utf8.txt", lambda p: p.write_text("AAA 評価\nBBB 電源\n耐久 試験\n", encoding="utf-8"))
    add(
        "text_cp932.txt",
        lambda p: p.write_bytes("電源ノイズ 測定 落下\n".encode("cp932")),
    )
    add(
        "text_utf16.txt",
        lambda p: p.write_text("温度 上昇 試験\n", encoding="utf-16"),
    )
    add("manual_2page.pdf", lambda p: write_text_pdf(p, ["AAA 評価", "BBB 耐久"]))
    if with_ocr_image:
        add("scan_page.pdf", lambda p: write_image_pdf(p, "AAA SCAN 1234"))
    add("report.xlsx", write_xlsx)
    add("review.docx", write_docx)
    add("slides.pptx", write_pptx)
    add("archive.zip", lambda p: write_zip(p, {"inner.txt": "異音 確認\n", "note.txt": "AAA 記録\n"}))
    add("unknown.dat", lambda p: p.write_text("落下 試験 データ\n", encoding="utf-8"))
    add("binary.bin", lambda p: p.write_bytes(os.urandom(4096)))

    # Comparison fixtures: the exact shapes the V4/V5 regression matrix needs.
    add(
        "parts.txt",
        lambda p: p.write_text("ABC-123 の記録\n2SC4117 と 2SC411T\nABC 123 の別表記\n", encoding="utf-8"),
    )
    add("nfkc.txt", lambda p: p.write_text("ＡＢＣ　ＤＥＦ の全角表記\n", encoding="utf-8"))
    add("case.txt", lambda p: p.write_text("aAa と BBB の混在\n", encoding="utf-8"))
    add("phrase_near.txt", lambda p: p.write_text("耐久 試験を実施\n電源とノイズの測定\n", encoding="utf-8"))

    nested = root_path / "sub" / "deep"
    nested.mkdir(parents=True, exist_ok=True)
    (nested / "nested_text.txt").write_text("耐久 サブフォルダ\n", encoding="utf-8")
    corpus.files["sub/deep/nested_text.txt"] = nested / "nested_text.txt"
    return corpus


def build_bench_corpus(root: str | os.PathLike[str], *, count: int = 2000) -> Corpus:
    """Larger corpus for benchmarks: mixed text/nested folders."""
    root_path = Path(root)
    root_path.mkdir(parents=True, exist_ok=True)
    corpus = Corpus(root=root_path)
    for index in range(count):
        folder = root_path / f"dir{index // 100:03d}"
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"doc{index:05d}.txt"
        body = [f"line {line} AAA 評価 {index}" for line in range(40)]
        if index % 7 == 0:
            body.append("BBB 電源")
        target.write_text("\n".join(body), encoding="utf-8")
        corpus.files[target.name] = target
    return corpus
