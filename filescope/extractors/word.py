"""Word extraction: python-docx plus the OOXML parts python-docx skips."""

from __future__ import annotations

from ..core.models import ChunkKind, FileEntry
from ..core.paths import WORD_EXTS
from ..errors import Issue, Severity
from .base import ExtractContext, ExtractOptions, ExtractResult, Sink
from .ooxml import iter_parts, open_zip, xml_texts


class WordExtractor:
    name = "word"
    version = "5.0"

    def supports(self, entry: FileEntry, options: ExtractOptions) -> bool:
        return entry.extension in WORD_EXTS

    def extract(self, context: ExtractContext, sink: Sink) -> ExtractResult:
        result = ExtractResult()
        limit = context.options.limits.office_max_bytes
        if context.entry.size > limit:
            result.skipped = True
            result.reason = f"Office上限（{limit // (1024 * 1024)}MB）を超えています"
            return result

        try:
            from docx import Document
        except ImportError:
            result.skipped = True
            result.reason = "python-docx が未導入です"
            result.warnings.append(
                Issue(path=context.entry.path, code="word-dependency", message=result.reason, severity=Severity.WARN)
            )
            return result

        try:
            document = Document(context.path)
        except Exception as exc:  # malformed or password protected
            result.skipped = True
            result.reason = f"開けませんでした: {type(exc).__name__}"
            result.warnings.append(
                Issue(path=context.entry.path, code="word-open", message=result.reason, severity=Severity.WARN)
            )
            return result

        for index, paragraph in enumerate(document.paragraphs, 1):
            sink.text(paragraph.text, ChunkKind.PARAGRAPH, f"段落 {index}")

        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                for column_index, cell in enumerate(row.cells, 1):
                    location = f"表{table_index} {row_index}行{column_index}列"
                    for paragraph in cell.paragraphs:
                        sink.text(paragraph.text, ChunkKind.PARAGRAPH, location)

        self._extract_parts(context, sink, result)
        return result

    def _extract_parts(self, context: ExtractContext, sink: Sink, result: ExtractResult) -> None:
        archive = open_zip(context.path)
        if archive is None:
            return
        kinds = {
            "header": (ChunkKind.HEADER, "ヘッダー"),
            "footer": (ChunkKind.FOOTER, "フッター"),
            "footnotes": (ChunkKind.FOOTNOTE, "脚注"),
            "endnotes": (ChunkKind.ENDNOTE, "文末脚注"),
            "comments": (ChunkKind.COMMENT, "コメント"),
        }
        with archive:
            for name, data in iter_parts(
                archive,
                lambda member: member.startswith("word/")
                and (
                    member.endswith("header1.xml")
                    or member.endswith("header2.xml")
                    or member.endswith("header3.xml")
                    or member.endswith("footer1.xml")
                    or member.endswith("footer2.xml")
                    or member.endswith("footer3.xml")
                    or member.endswith("footnotes.xml")
                    or member.endswith("endnotes.xml")
                    or member.endswith("comments.xml")
                    or "/header" in member
                    or "/footer" in member
                ),
            ):
                label = _label_for(name, kinds)
                for value in xml_texts(data, {"t"}):
                    sink.text(value, label[0], label[1])

            # Text boxes live inside the document part but outside paragraphs.
            for _name, data in iter_parts(archive, lambda member: member.endswith("document.xml")):
                for value in _textbox_texts(data):
                    sink.text(value, ChunkKind.TEXTBOX, "テキストボックス")


def _label_for(name: str, kinds: dict[str, tuple[ChunkKind, str]]) -> tuple[ChunkKind, str]:
    base = name.rsplit("/", 1)[-1]
    stem = base.split(".")[0].rstrip("0123456789")
    if stem in kinds:
        return kinds[stem]
    return ChunkKind.INTERNAL, "内部テキスト"


def _textbox_texts(data: bytes) -> list[str]:
    """Text inside ``<w:txbxContent>`` which python-docx does not expose."""
    import re

    out: list[str] = []
    for block in re.findall(rb"<w:txbxContent>(.*?)</w:txbxContent>", data, flags=re.DOTALL):
        for match in re.findall(rb"<w:t[^>]*>(.*?)</w:t>", block, flags=re.DOTALL):
            value = match.decode("utf-8", "replace").strip()
            if value:
                out.append(value)
    return out
